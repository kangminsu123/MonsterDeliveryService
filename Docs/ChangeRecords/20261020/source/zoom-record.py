from pathlib import Path
import json,shutil,hashlib
from docx import Document
from docx.shared import Inches,Pt
R=Path('C:/Users/xlfhf/Documents/Codex/2026-09-18/https-github-com-kangminsu123-monsterdeliveryservice-git/MonsterDeliveryService');D=R/'Docs/ChangeRecords/20261020';W=Path('C:/Users/xlfhf/Documents/Codex/2026-10-20/durl')
data={'title':'우클릭 유지 중 줌인',
'problem':'사용자는 우클릭을 누르는 동안 화면을 확대하는 기능을 요청했다. 기존에는 러닝화·배터리 이동 효과의 시야각 변화만 있었고 우클릭 줌인은 없었다.',
'method':'오른쪽 마우스 버튼의 isPressed 상태로 확대를 결정한다. 정상 시야각60도에서 누르는 동안35도로 부드럽게 전환하고, 떼면 러닝화·배터리 이동 상태에 맞는 현재 시야로 복귀한다. 로컬과 온라인 카메라 경로는 같은 UpdateCameraFieldOfView를 사용한다. 기존 손 아이템 카메라도 월드 카메라와 같은 시야각을 사용한다. 일시정지·네트워크 메뉴·경기 종료·맵 편집·커서 해제 상태에서는 줌을 적용하지 않는다. HUD 오른쪽에 우클릭 유지 줌인 안내를 추가했다. 이동속도·아이템 소모·판정 사거리와 프로토콜14는 변경하지 않았다.',
'verification':'2026년10월20일 Unity6000.5.10f1 SampleScene 실제 Editor Play에서 InputSystem 마우스 상태를 주입한15항목이 모두 통과했다. 우클릭 유지35도·해제60도 복귀, 빈손 사용, 손 카메라 시야각 일치, 아이템 횟수 미소모, 이동효과와 우선순위/해제 후 복귀,3인칭과 온라인 카메라 경로, 일시정지·메뉴·경기종료·편집모드·커서해제 제한을 확인했다. 최종 Console오류0. 같은 카메라 위치의 일반/줌 화면을 실제 Camera.Render로 촬영했다.',
'limits':'온라인 카메라 경로는 통제된 플레이어 레코드로 한Editor에서 검사했다. 별도 두프로세스 연결·WAN·배포 실행 파일 재빌드는 이번 미실행이다. 줌은 로컬 카메라 기능으로 네트워크 메시지나 프로토콜을 바꾸지 않았다. 사진의 위치와 입력은 검수용으로 설정했다.',
'retry':'첫 검사 코드에서 Mouse(bool) 보조 함수 이름이 InputSystem.Mouse 타입과 겹쳐 CS0119가 발생했다. 타입 이름을 완전히 지정한 뒤 컴파일 오류0 및15항목 통과를 확인했다.'}
(D/'source/zoom_report_data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
md='\n\n## 사례04 '+data['title']+'\n\n2026년10월20일\n'
for label,key in [('요청과 원인','problem'),('수정 방법','method'),('검증 결과','verification'),('확인 범위','limits'),('재시도','retry')]:md+='\n### '+label+'\n\n'+data[key]+'\n'
for f,caption in [('20261020_zoom_normal.png','일반 시야60도'),('20261020_zoom_held.png','같은 위치에서 우클릭 유지35도')]:md+='\n!['+caption+'](images/'+f+')\n\n'+caption+'\n'
md+='\n[실제Play15항목](evidence/hold-zoom-play.txt) · 메뉴 Tools/Gameplay/Camera/Verify Hold Zoom\n'
p=D/'README.md';p.write_text(p.read_text(encoding='utf8')+md,encoding='utf8')
p=D/'requests.csv';p.write_text(p.read_text(encoding='utf-8-sig')+'R09,우클릭유지중줌인,35도확대와해제시기존시야복귀,사례04 Play15항목\n',encoding='utf-8-sig')
doc=Document(D/'ChangeReport_20261020.docx');doc.add_page_break();doc.add_heading(data['title'],1)
doc.add_paragraph('2026년10월20일 우클릭을 누르는 동안 확대하고 떼면 원래 시야로 돌아오는 기능을 추가했다. 실제 입력 상태 검사15항목을 통과했으며 기존 게임 판정과 소모 규칙을 유지한다.')
for label,key in [('요청과 원인','problem'),('수정 방법','method'),('검증 결과','verification'),('확인 범위와 재시도','limits')]:doc.add_heading(label,1);doc.add_paragraph(data[key])
doc.add_paragraph(data['retry'])
doc.add_page_break();doc.add_heading('일반 화면과 우클릭 확대 화면',1)
for f,caption in [('20261020_zoom_normal.png','실제 Unity 일반 화면 시야각60도'),('20261020_zoom_held.png','같은 위치에서 우클릭 유지 시야각35도')]:doc.add_picture(str(D/'images'/f),width=Inches(6.2));doc.add_paragraph(caption,'Caption')
doc.add_paragraph('검사 원문 evidence/hold-zoom-play.txt. 같은 위치에서 실제 카메라를 렌더했다. 이미지에는 게임 HUD가 포함되지 않는다.')
doc.save(D/'ChangeReport_20261020.docx');shutil.copy2(D/'ChangeReport_20261020.docx',W/'outputs/ChangeReport_20261020.docx')
for f in ['20261020_zoom_normal.png','20261020_zoom_held.png']:shutil.copy2(D/'images'/f,W/'outputs'/f)
section='<section><h2>'+data['title']+'</h2><p>'+data['method']+'</p><p>'+data['verification']+'</p>'
for f in ['20261020_zoom_normal.png','20261020_zoom_held.png']:section+='<img src="images/'+f+'">'
with (D/'index.html').open('a',encoding='utf8') as stream:stream.write(section+'</section>')
note='## 우클릭 유지 줌인 2026년10월20일\n\n우클릭유지35도·해제시일반60도/이동효과시야복귀.로컬/온라인공통카메라경로·손카메라일치.메뉴/일시정지/종료/맵편집/커서해제시억제.HUD안내추가.Play15항목통과,Console0.프로토콜14유지.별도참가자/WAN/게임재빌드미실행.기록Docs/ChangeRecords/20261020/README.md사례04.\n\n'
for f in ['GAME_DESIGN.md','TEAM_STATUS.md','NEXT_SESSION.txt','ART_REVIEW.md']:
 p=R/f;p.write_text(note+p.read_text(encoding='utf-8-sig'),encoding='utf8')
files=['Assets/Gameplay/MonsterDeliveryPrototype.cs','Assets/Gameplay/DeliveryNetwork.cs','Assets/Gameplay/MonsterDeliveryHUD.cs','Assets/Editor/HoldZoomCheck.cs']
snapshot=D/'source/zoom-after';snapshot.mkdir(exist_ok=True)
for f in files:shutil.copy2(R/f,snapshot/Path(f).name)
(D/'evidence/zoom-code-sha256.json').write_text(json.dumps({f:hashlib.sha256((R/f).read_bytes()).hexdigest() for f in files},indent=2),encoding='utf8')
shutil.copy2(Path(__file__),D/'source/zoom-record.py')
print('Zoom case recorded with two actual captures')
