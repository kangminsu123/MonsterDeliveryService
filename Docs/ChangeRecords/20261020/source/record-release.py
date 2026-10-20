from pathlib import Path
from docx import Document
from docx.shared import Inches
import json,shutil,html,hashlib
W=Path('C:/Users/xlfhf/Documents/Codex/2026-10-20/durl');R=Path('C:/Users/xlfhf/Documents/Codex/2026-09-18/https-github-com-kangminsu123-monsterdeliveryservice-git/MonsterDeliveryService');D=R/'Docs/ChangeRecords/20261020'
logs={r:(D/f'evidence/park-connected/{r}.txt').read_text(encoding='utf8') for r in ['host','client']}
assert all('FAIL ' not in s for s in logs.values())
counts={r:sum(s.startswith('PASS ') for s in text.splitlines()) for r,text in logs.items()}
packages=json.loads((D/'evidence/package-release-verification.json').read_text(encoding='utf8'))
smoke=json.loads((D/'evidence/windows-release-smoke.json').read_text(encoding='utf8'))
data={
'title':'실제 멀티 검증과 윈도우 및 맥 실행 패키지',
'problem':'사용자는5분 경기와 공원 정기 택배를 멀티에서도 확인하고 완료한 내용을 자료실 HTML에 반영한 뒤 Windows와 macOS 실행 파일을 각각 패키징하도록 요청했다. 이전 단계의 검증은 한 Editor 내 통제된 프레임 적용이어서 별도 참가자의 실제 전송과 표시는 확인하지 않은 상태였다.',
'method':'NGO/UTP로 본 Unity Editor와 별도 읽기 전용 참가자 Editor를 localhost에 연결했다. 편도60ms 인위 지연과 매5번째 입력/프레임 전송 누락을 적용했다. 검수 도구가 방장의 남은 시간을 변경해60·120·180·240초와300초 종료를 재현했다. 양쪽 관찰값과 CSV를 별도로 기록했다. 공유 stage 파일은 입력 타이밍과 기대 위치만 전달하며 게임 상태는 각 프로세스가 실제 네트워크로 받은 값을 검사했다.',
'verification':f'최종 실제 연결 검사에서 방장{counts["host"]}항목과 참가자{counts["client"]}항목이 모두 통과했다. 초기5분·빈손·고정 테스트 무기 숨김, 첫 상자의 좌표·도착 안내·작은 빛, 분간 위치 유지와 최대1개, 참가자의 실제 E 입력에 의한 보상1개와 수집 후10초 이상 미출현,2·3·4분 출현 위치, 양쪽 우클릭35도와 해제60도 복귀를 확인했다. 조끼 면역·배터리 시간과 참가자에게 보이는 조끼 금빛 링 및 배터리 전기 아크, 경기 종료와 재경기 초기화, 실제 전송 누락 발생도 확인했다. 검사 중 런타임 및 종료 오류는 양쪽0이다.',
'retry':'첫 실행은 복제 프로젝트 Packages가 비어 있어 초기화되지 않아 캐시 연결을 복원했다. 복제 Editor 시작 시 기존 Unity Pipeline Unsafe DLL 진단도 나타났으며 검사 시작 전 환경 진단으로 구분했다. 첫 두 실제 연결 시도는 숨김 MPPM 입력 브리지가 가상 키와 마우스 이벤트를 걸러 수집과 참가자 줌 검사가 실패했다. 검수 도구 소유의 가상 장치 상태를 직접 설정하고 같은 게임 TickNetwork와 입력 처리 함수를 호출하도록 보강했다. 상자를 바라보는12방향 중 실제 광선이 닿는 접근점도 선택했다. 게임 런타임의 수집·줌 로직은 바꾸지 않았다. 최초 실패 로그와 CSV를 보존했으며 마지막 실제 연결 검사는 전 항목 통과했다.',
'builds':'설치된 Unity6000.5.10f1 플랫폼 모듈로 같은 게임 소스와Protocol15의 Mono 일반 빌드를 만들었다. Windows x64와 macOS Universal(Intel x86_64 및 Apple Silicon arm64)을 각각 LZ4HC로 빌드했다. 각 패키지에 사용 안내와 한 PC에서2명/4명 연결하는 실행 도우미, 소스와 어셈블리 SHA256을 담은 BUILD_INFO.json을 포함했다. ZIP CRC와 파일 구조를 검사했으며 맥 Mach-O 바이너리의 두 아키텍처와 ZIP 내 실행 권한도 확인했다. 두 빌드 결과는Succeeded다. Windows 보고서에오류1·경고4, Mac에오류4·경고13이 기록됐다. 기존 Pipeline DLL 및 Mac 임포트 때 사용하지 않는 ShaderGraph BuiltIn 예제 템플릿의 경로 읽기 진단이다. 해당 예제의 게임Assets 참조가 없는 점을 확인했으며 원문 진단과 성공 결과를 함께 보존했다.',
'smoke':'새 Windows 실행 파일2개를 숨김으로 실행해 호스트와 참가자의 실제 연결 로그 및 런타임 오류 유무를 확인했다. '+smoke.get('summary',''),
'library':'웹 자료실과 바탕화면 MonsterDeliveryLibrary.html에 현재 경기 기준·우클릭 줌·실제 멀티 결과·Windows/Mac 빌드 설명을 추가한다. 기존 아이템 이미지와 모델을 유지한다. HTML에는 최신 보고서와 실제 캡처를 내장하며 대용량 실행 ZIP은 옆의 MonsterDeliveryBuilds_20261020 폴더를 통해 연결한다. 웹 자료실은 빌드 안내와 검증 기록을 제공하고 ZIP은 별도로 전달한다.',
'limits':'실제5분을 기다린 시험은 아니며 시간 경계를 이동해 검증했다. 두 프로세스는 같은 Windows PC의 localhost이고 실제WAN·다른 물리PC·4인전은 미실행이다. macOS 앱은 구조와 아키텍처를 확인했으며 Mac 실기 실행 및 Windows↔Mac 실제 접속은 이 Windows 호스트에서 실행하지 않았다. Apple Developer ID 서명과 공증도 수행하지 않았다. 사용자 Unity 저장소는 커밋하거나 푸시하지 않았다.'}
(D/'source/release_report_data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
shutil.copy2(D/'ChangeReport_20261020.docx',D/'source/release_before_report.docx')
md='\n\n## 사례06 '+data['title']+'\n\n작성 날짜2026년10월20일 · 요청R15~R18 · 실제 연결 확인과 배포 산출물\n'
for key,label in [('problem','요청'),('method','검사 방법'),('verification','결과'),('retry','최초 실패와 보강'),('builds','실행 패키지'),('smoke','Windows 실행 확인'),('library','자료실 갱신'),('limits','확인 범위')]:md+='\n### '+label+'\n\n'+data[key]+'\n'
for platform,p in packages.items():md+=f'\n{platform} ZIP {p["bytes"]:,}bytes · SHA256 {p["sha256"]}\n'
for r in ['host','client']:md+=f'\n![실제 {r} 상자 화면](images/20261020_park_connected_{r}.png)\n\n2026년10월20일 별도 프로세스의 실제 카메라 렌더. 검수용 위치를 지정했다.\n'
md+='\n[방장 검사](evidence/park-connected/host.txt) · [참가자 검사](evidence/park-connected/client.txt) · [패키지 검증](evidence/package-release-verification.json) · [Windows 실행 확인](evidence/windows-release-smoke.json)\n'
p=D/'README.md';p.write_text(p.read_text(encoding='utf8')+md,encoding='utf8')
p=D/'requests.csv';s=p.read_text(encoding='utf-8-sig')
for row in [('R15','실제 멀티','별도 프로세스 연결과 출현 수집 줌 연출'),('R16','자료실 HTML','완료 내용과 보고서 실제 캡처 추가'),('R17','Windows 패키지','x64 빌드와 ZIP 및 실행 연결 확인'),('R18','macOS 패키지','Universal 빌드와 실행 권한 보존 ZIP')]:s+=','.join(row)+',사례06,검증 및 패키징 완료\n'
p.write_text(s,encoding='utf-8-sig')
doc=Document(D/'ChangeReport_20261020.docx');doc.add_page_break();doc.add_heading(data['title'],1)
doc.add_paragraph(f'별도 두 프로세스의 실제 연결에서 방장{counts["host"]}항목과 참가자{counts["client"]}항목이 통과했다. 같은Protocol15 소스로 Windows x64와 macOS Universal 실행 패키지를 생성하고 파일 무결성을 확인했다.')
for label,key in [('요청과 검사 방법','method'),('실제 연결 결과','verification'),('재시도와 검사 환경','retry'),('남은 확인 범위','limits')]:doc.add_heading(label,1);doc.add_paragraph(data[key])
doc.add_page_break();doc.add_heading('플랫폼별 패키지와 확인 범위',1)
doc.add_paragraph(data['builds'])
for platform,p in packages.items():doc.add_heading(platform+' 실행 파일',1);doc.add_paragraph(p['file']+'\n'+f'{p["bytes"]:,}bytes · ZIP CRC 통과\nSHA256 '+p['sha256'])
for label,key in [('Windows 실제 실행','smoke'),('자료실 반영','library')]:doc.add_heading(label,1);doc.add_paragraph(data[key])
doc.add_page_break();doc.add_heading('두 프로세스에서 보이는 공원 택배',1)
for r,label in [('host','실제 방장 화면'),('client','실제 참가자 화면')]:doc.add_picture(str(D/f'images/20261020_park_connected_{r}.png'),width=Inches(6.2));doc.add_paragraph(label,'Caption')
doc.add_paragraph('2026년10월20일 실제 Unity 카메라를 검수용 위치에 두고 렌더했다. 화면만으로 수집이나 동기화 판정을 주장하지 않으며 각 프로세스의 검사 로그와 시간별 CSV로 함께 확인했다.')
doc.save(D/'ChangeReport_20261020.docx');shutil.copy2(D/'ChangeReport_20261020.docx',W/'outputs/ChangeReport_20261020.docx')
section='<section><h2>'+data['title']+'</h2>'+''.join('<p>'+html.escape(v)+'</p>' for k,v in data.items() if k!='title')
for r in ['host','client']:section+=f'<img src="images/20261020_park_connected_{r}.png">'
with (D/'index.html').open('a',encoding='utf8') as f:f.write(section+'</section>')
note=f'## 실제 멀티와 양 플랫폼 패키징 2026년10월20일\n\n두 Editor 실제localhost·편도60ms·20퍼센트 입력/상태누락.방장{counts["host"]}·참가자{counts["client"]}항목통과,검사중/종료오류0.출현좌표/안내/원격E수집/보상1개/매분/종료/재경기·양쪽줌·상대조끼와배터리표시확인.Protocol15.Windows x64와Mac Universal Mono일반빌드·ZIP CRC·Mac 실행권한검사.새Windows2개실행연결확인.Mac실기/Windows↔Mac/WAN/4인미검증.자료실HTML과웹최신내용반영은evidence/library-release-final.json참조.보고서사례06.게임커밋푸시없음.\n\n'
for f in ['NEXT_SESSION.txt','GAME_DESIGN.md','TEAM_STATUS.md','ART_REVIEW.md']:
 p=R/f;p.write_text(note+p.read_text(encoding='utf-8-sig'),encoding='utf8')
snap=D/'source/release-after';snap.mkdir(exist_ok=True)
for f in ['Assets/Editor/ParkDeliveryNetworkCheck.cs','Assets/Editor/DeliveryReleaseBuild.cs']:shutil.copy2(R/f,snap/Path(f).name)
shutil.copy2(__file__,D/'source/record-release.py')
print(json.dumps({'host':counts['host'],'client':counts['client'],'case':6,'report':'appended'},ensure_ascii=False))
