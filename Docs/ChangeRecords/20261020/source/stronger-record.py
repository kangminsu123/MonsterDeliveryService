from pathlib import Path
import json,shutil,html
from docx import Document
R=Path('C:/Users/xlfhf/Documents/Codex/2026-09-18/https-github-com-kangminsu123-monsterdeliveryservice-git/MonsterDeliveryService');D=R/'Docs/ChangeRecords/20261020'
data={'title':'원격 배송기 사거리 두 배와 취급주의 광채 강화',
'problem':'사용자가 특급 배송기 사거리를 약 두 배 늘리고 약한 이펙트를 더 강하게 해 달라고 요청했다. 우체통 점령용 원격 배송기로 해석했음을 작업 시작 때 알렸다.',
'method':'RemoteDeliveryRange 공통 상수를 36m로 두어 기존18m의 두 배로 확장했다. 로컬·서버의 RemoteCaptureTarget, 조준 대상 검색과 발사선이 같은 값을 쓴다. 실패 안내도36m로 바꿨다. 발사선 두께0.08→0.2m. 취급주의 택배는 링 두께0.085→0.18m, 광원 범위3.2→5m, 밝기0.9~1.4→3~5로 강화했다. 반짝임6→12개, 크기와 두께를 키우고 택배 주변에1.9m 높이의 회전 상승 나선을 추가했다. E 수집·수집 시 숨김·재출현 복구·아군 제외·벽 차단·20초 점령 보호·1회 소모는 유지한다.',
'verification':'실제 Unity6000.5.10f1 SampleScene Editor Play에서 RemoteDeliveryCheck20항목 통과. 35m 점령 성공·37m 미소모, 서버actor0/1 각각30m 점령과 아군 차단, 벽·빈 조준·기절 제한을 확인했다. ParkDeliveryCheck33항목도 통과해 새26개 LineRenderer, 실제 시간 애니메이션, 광원5m·밝기상한5, 충돌체 추가 없음, E수집·숨김·재출현·호스트·스냅샷을 확인했다. 최종 Console오류0. 테스트 후 Play 종료. 검사군은 중복이 있으므로 고유 사례 수로 합산하지 않는다.',
'limits':'별도 두 프로세스 연결·WAN·배포 실행파일 재빌드 미실행. 프로토콜15와 메시지 형식 유지. 이미 배포된 실행파일에는 미포함이다. 캡처는 실제 Unity에서 시계·위치·카메라를 설정한 검수용 재현이다. Word 본문은 갱신하되 기존 LibreOffice 미설치로 페이지 렌더 검수는 미완료다.'}
(D/'source/stronger_report_data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
md='\n\n## 사례09 '+data['title']+'\n\n2026-10-20 · 요청R24 R25 · Editor 확인 완료\n\n'+'\n\n'.join(data[k] for k in ['problem','method','verification','limits'])+'\n\n![강화된 금빛 나선과 반짝임](images/20261020_fragile_strong_notice.png)\n\n[원격배송20항목](evidence/remote-delivery-36m-play.txt) · [광채33항목](evidence/park-delivery-strong-play.txt)\n'
p=D/'README.md';p.write_text(p.read_text(encoding='utf8')+md,encoding='utf8')
p=D/'requests.csv';p.write_text(p.read_text(encoding='utf-8-sig')+'R24,원격 배송기,사거리18m에서36m로 확장,사례09,Editor20항목 통과\nR25,이펙트,택배 광채와 배송기 발사선 강화,사례09,Editor33항목 통과\n',encoding='utf-8-sig')
p=D/'index.html';p.write_text(p.read_text(encoding='utf8')+'<section><h2>'+data['title']+'</h2>'+''.join('<p>'+html.escape(data[k])+'</p>' for k in ['problem','method','verification','limits'])+'<img src="images/20261020_fragile_strong_notice.png" alt="실제Unity 강화 효과"></section>',encoding='utf8')
doc=Document(D/'ChangeReport_20261020.docx');doc.add_page_break();doc.add_heading(data['title'],1)
for k in ['problem','method','verification','limits']:doc.add_paragraph(data[k])
doc.save(D/'ChangeReport_20261020.docx')
note='## 원격 배송기36m와 이펙트 강화 2026년10월20일\n\n원격배송기18→36m·조준/서버/발사선 일치, 선폭0.2m. 취급주의 금빛 링 강화·12반짝임·상승나선·광원5m 밝기3~5. 원격20·공원33항목통과,Console0. 실제 캡처 및 기록 Docs/ChangeRecords/20261020/README.md 사례09. 별도참가자/WAN/재빌드 미실행. Word 페이지검수는 기존LibreOffice 미설치로 미완료.\n\n'
for f in ['NEXT_SESSION.txt','GAME_DESIGN.md','TEAM_STATUS.md','ART_REVIEW.md']:
 p=R/f;p.write_text(note+p.read_text(encoding='utf8'),encoding='utf8')
snap=D/'source/stronger-after';snap.mkdir(exist_ok=True)
for f in ['Assets/Gameplay/SpecialDeliveryItems.cs','Assets/Gameplay/MonsterDeliveryPrototype.cs','Assets/Gameplay/ParkDeliveryGlow.cs','Assets/Editor/RemoteDeliveryCheck.cs','Assets/Editor/ParkDeliveryCheck.cs','Assets/Editor/ParkDeliveryNetworkCheck.cs']:shutil.copy2(R/f,snap/Path(f).name)
shutil.copy2(__file__,D/'source/stronger-record.py')
print('case09 recorded')
