from pathlib import Path
import json,shutil,hashlib,html
from docx import Document
from docx.shared import Inches
R=Path('C:/Users/xlfhf/Documents/Codex/2026-09-18/https-github-com-kangminsu123-monsterdeliveryservice-git/MonsterDeliveryService');D=R/'Docs/ChangeRecords/20261020';W=Path('C:/Users/xlfhf/Documents/Codex/2026-10-20/durl')
data={
'title':'5분 경기와 매분 공원 취급주의 택배',
'problem':'사용자는 경기 시간을5분으로 늘리고 매분 도착 안내와 함께 공원 내 무작위 위치에 취급주의 상자1개가 출현하도록 요청했다. 기존 경기 기본값은180초였으며 상자는 개발용 고정 위치에서10초마다 재출현했다. 공원 고정 무기5종과 시작 장비3종도 개발용 규칙이었다.',
'method':'경기 기본값을300초로 바꿨다. 실제 경기는 시작 후60·120·180·240초에 한 번씩 도착하며 종료 시점300초에는 추가 출현하지 않는다. 이전 상자가 남아 있으면 새 위치로 옮겨 최대1개만 유지한다. 공원의 실제 석재 광장과 진입로 메시 위에서 무작위로 위치를 뽑고 네 모서리 지면 및 상자 주변 장애물 겹침을 검사한다. 수집한 상자는 다음 정각까지 돌아오지 않는다. 실제 경기에서는 고정 테스트 무기와 시작 지급을 숨기고 개발용 혼자 탐색은 기존 테스트 구성을 유지한다.',
'presentation':'출현할 때 공원에 취급주의 택배가 도착했습니다! 문구를5초간 표시한다. 기존 상자 모델에 낮은 금빛 링1개와 반경2.2m의 그림자 없는 광원1개를 추가했다. 밝기는0.35~0.5로 천천히 변하며 수집 즉시 모델·빛·충돌체가 함께 숨겨진다. 별도 거대 기둥이나 강한 섬광은 없다.',
'rewards':'취급주의 택배 보상은 공기 대포·긴급 배송 방송기·원격 배송기·구명조끼·특급 배송 배터리5종이며 각각20%다. 기존 Random.Range(10,15) 균등 추첨을 재사용했다. 매 수집마다 방장이1종을 결정하며 기존 사용 횟수와 아이템 효과는 유지한다. 일반 택배의1~9번 아이템 추첨과100후보·50상자 보충 규칙은 변경하지 않았다.',
'network':'온라인에서는 방장만 경기 시계로 출현·수집·보상을 확정한다. 기존 스냅샷에 상자 위치·가용 상태·도착 분·안내 남은 시간을 추가해 참가자가 동일하게 표시한다. 새 가용 상태가 자동 재출현 타이머로 해석되지 않도록 분리했다. 프로토콜14에서15로 올려 이전 클라이언트의 혼합 접속을 막았다.',
'verification':'2026년10월20일 Unity6000.5.10f1 SampleScene 실제 Play31항목 통과, 최종 Console오류0. 05:00 시작, 60초 직전과직후·같은분 중복·수집 후71초와119초·120/180/240초·300초 종료·재시작을 확인했다. 공원 무작위 위치120개와 보상10000회 추첨을 확인했다. 실제 NGO 호스트의 경기 시작·출현·수집 광선·보상1개 및 JSON 프레임 왕복 후 위치·표시·수집 동기화 경로를 검사했다. 추첨 횟수는 순서대로2054·2029·2003·1926·1988이었다.',
'limits':'검사 도구가 경기 시작 시각과 종료 시각을 이동시켜 시간 경계를 재현했다. 실제5분을 끝까지 기다린 시험은 아니다. NGO 호스트 및 참가자 프레임 적용은 한 Editor에서 검사했으며 별도 두 프로세스 참가자·WAN·릴리스 실행 파일 재빌드는 미실행이다. 정지된 카메라에서 실제 HUD 그리기 함수를 호출한 캡처를 포함한다. 웹 자료실과 바탕화면 HTML은 이번 요청 범위에서 갱신하지 않았다.',
'retry':'첫 검사에서 광원 생성 직후 기본 밝기가 한 프레임 남아28항목 중1항목이 실패했다. Awake에서 초기 색과 밝기를 즉시 지정해 수정했다. 화면 검수에서 숨긴 테스트 무기의 이름표가 남는 점도 찾아 비활성 오브젝트의 이름표를 그리지 않도록 고쳤다. 카메라가 벤치에 가린 캡처와 HUD 없는 첫 캡처를 폐기하고 원본 HUD 함수와 고정 검수 카메라로 재촬영했다. 최종31항목 모두 통과했다.'}
source=D/'source';source.mkdir(exist_ok=True)
(source/'park_delivery_report_data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
shutil.copy2(D/'ChangeReport_20261020.docx',source/'park_delivery_before_report.docx')
md='\n\n## 사례05 '+data['title']+'\n\n작성 날짜2026년10월20일 · 요청R10~R14 · 요구 변경과 연출 개선 · 상태 Editor 확인 완료\n'
for label,key in [('요청과 원인','problem'),('수정 방법','method'),('연출','presentation'),('랜덤 보상','rewards'),('온라인 처리','network'),('검증 결과','verification'),('확인 범위','limits'),('재시도','retry')]:md+='\n### '+label+'\n\n'+data[key]+'\n'
for f,c in [('20261020_park_delivery_notice.png','실제 도착 안내와 남은 시간04:00 및 상자 빛'),('20261020_park_delivery_glow.png','실제 상자 모델과 은은한 링의 근접 화면')]:md+='\n!['+c+'](images/'+f+')\n\n2026년10월20일 검수용 위치와 시각을 설정한 실제 Unity 렌더. '+c+'\n'
md+='\n메뉴 Tools/Gameplay/Special Delivery/Verify Minute Park Delivery · [최종31항목](evidence/park-delivery-play.txt) · [첫 검사](evidence/park-delivery-play-first-attempt.txt) · [JSON 프레임](evidence/park-delivery-frame.json)\n'
p=D/'README.md';p.write_text(p.read_text(encoding='utf8')+md,encoding='utf8')
p=D/'requests.csv';s=p.read_text(encoding='utf-8-sig')
for row in [('R10','경기 시간','기본 경기300초와 종료 확인'),('R11','도착 안내','매분 안내5초 표시'),('R12','상자 리스폰','공원 랜덤 지점에 최대1개'),('R13','상자 이펙트','은은한 금빛 링과 작은 광원'),('R14','랜덤 무기','특수5종 균등20퍼센트')]:s+=','.join(row)+',사례05,Editor 확인 완료\n'
p.write_text(s,encoding='utf-8-sig')
doc=Document(D/'ChangeReport_20261020.docx');doc.add_page_break();doc.add_heading(data['title'],1)
doc.add_paragraph('경기를5분으로 늘리고 매분 공원에 취급주의 택배1개가 도착하도록 구현했다. 출현 안내와 은은한 빛을 추가했으며5종 보상을 각각20%로 추첨한다. 실제 Editor31항목을 통과했다.')
for label,key in [('요청과 원인','problem'),('출현과 수집','method'),('빛과 보상','presentation')]:doc.add_heading(label,1);doc.add_paragraph(data[key])
doc.add_paragraph(data['rewards'])
doc.add_page_break();doc.add_heading('온라인 처리와 확인 근거',1)
for label,key in [('온라인 처리','network'),('실제 Editor 검증','verification'),('검사 범위','limits'),('최초 실패와 수정','retry')]:doc.add_heading(label,1);doc.add_paragraph(data[key])
doc.add_paragraph('검사 원문 evidence/park-delivery-play.txt. 메뉴 Tools/Gameplay/Special Delivery/Verify Minute Park Delivery. 최초 검사와 JSON 프레임도 같은 evidence 폴더에 보존했다.')
doc.add_page_break();doc.add_heading('실제 출현 안내와 상자 연출',1)
for f,c in [('20261020_park_delivery_notice.png','04:00에 도착 안내와 함께 나타난 상자'),('20261020_park_delivery_glow.png','약한 금빛 링과 근접 상자 화면')]:doc.add_picture(str(D/'images'/f),width=Inches(6.2));doc.add_paragraph(c,'Caption')
doc.add_paragraph('2026년10월20일 실제 Unity 캡처. 검사 도구로 시각과 카메라 위치를 지정했다. 안내 화면은 시뮬레이션을 멈추고 실제 HUD 그리기 함수를 호출해 촬영했으며 문구를 이미지에 합성하지 않았다.')
doc.save(D/'ChangeReport_20261020.docx')
shutil.copy2(D/'ChangeReport_20261020.docx',W/'outputs/ChangeReport_20261020.docx')
section='<section><h2>'+data['title']+'</h2>'
for key in data:
 if key!='title':section+='<p>'+html.escape(data[key])+'</p>'
for f in ['20261020_park_delivery_notice.png','20261020_park_delivery_glow.png']:
 section+='<img src="images/'+f+'">';shutil.copy2(D/'images'/f,W/'outputs'/f)
with (D/'index.html').open('a',encoding='utf8') as out:out.write(section+'</section>')
note='## 5분 경기와 매분 공원 택배 2026년10월20일\n\n경기300초.60·120·180·240초에공원석재광장과진입로무작위안전지점에취급주의택배최대1개.도착안내5초·은은한금빛링과광원.보상5종각20퍼센트.실제경기테스트시작장비·고정5종숨김,혼자탐색보존.방장권한·위치/가용상태/안내스냅샷추가·프로토콜15.실제Editor31항목통과Console오류0.경기시계이동검수,별도참가자/WAN/실행파일재빌드미실행.자료실/바탕화면HTML이번미변경.기록Docs/ChangeRecords/20261020/README.md사례05.아래180초·프로토콜14·경기10초리스폰은이전기록.\n\n'
for f in ['GAME_DESIGN.md','TEAM_STATUS.md','ART_REVIEW.md','NEXT_SESSION.txt']:
 p=R/f;p.write_text(note+p.read_text(encoding='utf-8-sig'),encoding='utf8')
files=['Assets/Gameplay/'+f for f in ['MonsterDeliveryPrototype.cs','SpecialDeliveryItems.cs','SpecialDeliveryParcel.cs','ParkDelivery.cs','ParkDeliveryGlow.cs','MonsterDeliveryHUD.cs','DeliveryNetwork.cs']]+['Assets/Editor/ParkDeliveryCheck.cs']
snap=source/'park-delivery-after';snap.mkdir(exist_ok=True)
hashes={}
for f in files:shutil.copy2(R/f,snap/Path(f).name);hashes[f]=hashlib.sha256((R/f).read_bytes()).hexdigest()
(D/'evidence/park-delivery-code-sha256.json').write_text(json.dumps(hashes,indent=2),encoding='utf8')
shutil.copy2(__file__,source/'park-delivery-record.py')
print('Updated case05 and 31-pass evidence; report appended three sections.')
