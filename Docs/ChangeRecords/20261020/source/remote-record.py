from pathlib import Path
import json,shutil,hashlib
from docx import Document
from docx.shared import Inches,Pt
R=Path('C:/Users/xlfhf/Documents/Codex/2026-09-18/https-github-com-kangminsu123-monsterdeliveryservice-git/MonsterDeliveryService');D=R/'Docs/ChangeRecords/20261020';W=Path('C:/Users/xlfhf/Documents/Codex/2026-10-20/durl');S=Path('C:/Users/xlfhf/Documents/Codex/2026-10-17/dbslx/work/monster-delivery-hub')
data={'title':'자동 포장기를 원격 배송기로 교체',
'problem':'사용자는 상대를 상자로 감싸 기절시키는 포장 기능을 제거하고, 기존 이미지를 유지한 채 원거리에서 우체통을 점령하고20초 재점령 보호도 무시하는 아이템으로 교체하기를 요청했다.',
'method':'표시 이름을 원격 배송기로 변경했다. 기존 PackingMachine 열거값·프리팹·FBX·아이콘과 노란 빔을 재사용해 이미지 파일과 모델은 변경하지 않았다.18m 안 조준선의 첫 충돌이 상대 또는 중립 우체통일 때 즉시 점령한다. 자기 또는 같은 팀 소유, 벽이나 캐릭터로 막힌 조준, 사거리 밖, 빗나간 조준에는 소모와 재사용 대기시간을 시작하지 않는다. 성공 시1회 소모하며 기존 짧은 사용 연출 뒤 슬롯을 비운다. 포장 상자와 기절 판정은 제거했다.',
'capture':'공통 CaptureMailbox에서 성공 후 소유자·색상·점수·새20초 보호를 적용한다. 일반 Capture는 기존20초 보호를 계속 지키며 원격 배송기만 보호 무시 경로를 사용한다. 로컬과 서버 명령에서 같은 조준 검사를 먼저 실행해 유효한 사용만 소모한다. 서버가 소유권을 확정하고 기존 소유자 및 보호 잔여 시간 스냅샷으로 참가자에게 전달한다. 동작이 다른 구버전과 혼합하지 않도록 프로토콜을13에서14로 올렸다.',
'verification':'실제 Unity6000.5.10f1 SampleScene Play 원격 배송기19항목 통과. 중립 즉시 점령, 보호 중 상대 점령, 새20초 보호,1회 소모·슬롯 정리, 내 우체통·사거리·벽·빈 조준·캐릭터·기절 제한, 캐릭터 포장/기절 없음, 서버 명령source0/1의 적·아군 판정, 실제 맵8m 조준·색상/점수 변화를 확인했다. 공원5종 회귀44항목 통과. 일반 점령 보호 검사는 별도 로그에서 결과를 확인한다.',
'retry':'첫 검수는 캐릭터 배치 높이가 맞지 않았고 이후 상태를 초기화하지 않아 두 검사가 실패했다. 실제 우체통 검수 카메라의 회전과 눈 위치를 맞춘 뒤 재실행해19항목을 통과했다. 첫 화면은 캡처 완료 전에 검수 상태를 복구해 잘못된 슬롯이 보였으므로 사용 전후 프레임 대기 후 다시 촬영했다. 원거리 HUD 문구 줄바꿈도 줄였다. 기존 보호 검사에는 옛프로토콜6 기대값과 시작3슬롯을 비우지 않아 장갑을 얻지 못하는 검수 조건이 있어 현재14와 빈슬롯 조건으로 갱신했다. 첫 로그를 별도 보존했다.',
'limits':'서버 명령과 참가자 스냅샷의 검수는 한Editor의 통제된 레코드 검사다. 이번 별도 네트워크 두프로세스·WAN·4인전과 Windows/Mac 배포 게임 재빌드는 수행하지 않았다. 최신 프로토콜14 소스를 모든 참가자가 사용해야 하며 기존Protocol13 실행 파일은 별도 재빌드가 필요하다. 이미지 속 플레이어 위치와 우체통 소유 상태는 검수용으로 설정한 실제 Unity 화면이다.'}
(D/'source/remote_report_data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
md='\n\n## 사례03 '+data['title']+'\n\n2026년10월20일\n'
for label,key in [('문제와 요청','problem'),('수정 방법','method'),('점령과 네트워크','capture'),('검증 결과','verification'),('재시도','retry'),('확인 범위','limits')]:md+='\n### '+label+'\n\n'+data[key]+'\n'
for f,caption in [('20261020_remote_delivery_before.png','보호 중 상대 우체통과 기존 모델'),('20261020_remote_delivery_capture.png','원격 점령 성공과1회 소모 및 새20초 보호')]:md+='\n!['+caption+'](images/'+f+')\n\n'+caption+'\n'
md+='\n[원격19항목](evidence/remote-delivery-play.txt) · [공원44항목](evidence/remote-row-regression.txt) · [일반 점령 보호](evidence/remote-capture-regression.txt)\n'
p=D/'README.md';p.write_text(p.read_text(encoding='utf8')+md,encoding='utf8')
p=D/'requests.csv';p.write_text(p.read_text(encoding='utf-8-sig')+'R07,포장을원격점령아이템으로교체,18m즉시점령과보호20초무시,사례03\nR08,기존이미지유지하고이름내용변경,원격배송기이름과기존모델아이콘재사용,사례03\n',encoding='utf-8-sig')
doc=Document(D/'ChangeReport_20261020.docx');doc.add_page_break();doc.add_heading(data['title'],1)
doc.add_paragraph('2026년10월20일 요구 변경이다. 포장 기능을 제거하고 기존 이미지·모델을 그대로 사용하는 원격 배송기로 바꿨다.18m 즉시점령과20초 보호 무시를 실제Play19항목으로 확인했다.')
for label,key in [('문제와 요청','problem'),('수정 방법','method'),('점령과 네트워크','capture'),('검증과 확인 범위','verification')]:doc.add_heading(label,1);doc.add_paragraph(data[key])
doc.add_page_break();doc.add_heading('원격 점령 실제 Unity 화면',1)
for f,caption in [('20261020_remote_delivery_before.png','보호 중 상대 우체통을 약8m에서 조준한 상태'),('20261020_remote_delivery_capture.png','좌클릭 뒤 점수0에서1로 변경되고 아이템1회 소모')]:doc.add_picture(str(D/'images'/f),width=Inches(6.2));doc.add_paragraph(caption,'Caption')
doc.add_page_break();doc.add_heading('원격 배송기 검사 근거',1)
doc.add_paragraph(data['limits']);doc.add_paragraph(data['retry'])
for line in (D/'evidence/remote-delivery-play.txt').read_text(encoding='utf-8-sig').splitlines():
 if line.startswith(('AUTHORED','RAY')):continue
 p=doc.add_paragraph(line);p.paragraph_format.space_after=Pt(2)
 for run in p.runs:run.font.size=Pt(9)
doc.save(D/'ChangeReport_20261020.docx')
for f in ['ChangeReport_20261020.docx']:shutil.copy2(D/f,W/'outputs'/f);shutil.copy2(D/f,S/'public/library'/f)
for f in ['20261020_remote_delivery_before.png','20261020_remote_delivery_capture.png']:shutil.copy2(D/'images'/f,W/'outputs'/f);shutil.copy2(D/'images'/f,S/'public/library'/f)
note='## 원격 배송기 교체 2026년10월20일\n\n기존포장기모델·아이콘유지.표시이름원격배송기,18m상대/중립우체통즉시점령·기존20초보호무시·성공1회소모·새20초보호.자기/아군·벽·범위밖·빈조준미소모.포장/기절제거.원격Play19·공원44통과.프로토콜14.별도실제참가자/WAN/게임재빌드미실행.기록Docs/ChangeRecords/20261020/README.md사례03.아래포장기와Protocol13은이전기록.\n\n'
for name in ['GAME_DESIGN.md','TEAM_STATUS.md','NEXT_SESSION.txt','ART_REVIEW.md']:
 p=R/name;p.write_text(note+p.read_text(encoding='utf-8-sig'),encoding='utf8')
files=['Assets/Gameplay/SpecialDeliveryItems.cs','Assets/Gameplay/MonsterDeliveryPrototype.cs','Assets/Gameplay/MonsterDeliveryHUD.cs','Assets/Gameplay/DeliveryNetworkGameplay.cs','Assets/Gameplay/DeliveryNetwork.cs','Assets/Editor/SpecialDeliverySetup.cs','Assets/Editor/SpecialItemRowCheck.cs','Assets/Editor/CaptureProtectionCheck.cs','Assets/Editor/RemoteDeliveryCheck.cs']
after=D/'source/remote-after';after.mkdir(exist_ok=True)
for f in files:shutil.copy2(R/f,after/Path(f).name)
(D/'evidence/remote-source-sha256.json').write_text(json.dumps({f:hashlib.sha256((R/f).read_bytes()).hexdigest() for f in files},indent=2),encoding='utf8')
assets=['Assets/Resources/UI/InGame/item_PackingMachine.png','Assets/Resources/DeliveryItems/PackingMachine.prefab','Assets/Art/SpecialDelivery/Models/PackingMachine.fbx']
(D/'evidence/remote-retained-assets.json').write_text(json.dumps({f:hashlib.sha256((R/f).read_bytes()).hexdigest() for f in assets},indent=2),encoding='utf8')
print('Remote-delivery case recorded')
