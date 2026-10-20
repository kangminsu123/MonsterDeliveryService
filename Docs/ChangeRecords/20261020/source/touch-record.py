from pathlib import Path
from docx import Document
from docx.shared import Inches
import json, shutil, html
R=Path('C:/Users/xlfhf/Documents/Codex/2026-09-18/https-github-com-kangminsu123-monsterdeliveryservice-git/MonsterDeliveryService')
D=R/'Docs/ChangeRecords/20261020'
data={
'title':'일반 택배 접촉 자동 획득과 취급주의 E 유지',
'problem':'일반 택배도 조준 후 E키 입력이 필요했다. 사용자는 일반 택배는 몸이 닿으면 자동 획득하고 취급주의 택배는 기존 E키 획득을 유지하도록 요청했다. 원인은 로컬 TryPickup과 서버 NetCaptureAndPickup의 조준 및 E 입력 조건이었다.',
'method':'FloatingParcel.CollectTouching에서 캐릭터 캡슐과 접촉 스킨 범위의 Collider를 확인하고 기존 TryCollect와 풀 점유 해제를 재사용한다. 로컬 이동 후 AutoPickupParcels를 실행하며 서버는 모든 연결된 플레이어의 이동 처리 후 E 상태와 별개로 접촉을 검사한다. 빈 슬롯에 기존 일반 9종 중 하나를 지급한다. 슬롯이 가득 찼거나 기절·행동 중·넉백·대시이면 수집하지 않는다. 취급주의와 공원 고정 아이템은 기존 E 경로를 유지한다. 일반 택배의 원거리 E 획득은 제거하고 HUD를 닿으면 자동 획득으로 바꿨다. 메시지 형식과 프로토콜15는 유지한다.',
'verification':'Unity 6000.5.10f1 SampleScene 실제 Editor Play에서 ParcelTouchCheck 19항목 모두 통과, Console 오류0. 조준·E 없는 접촉 획득, 먼 거리 E 차단, 중복 방지, 슬롯 가득 참과 슬롯 비운 후 획득, 기절·대시·비권한 차단, 취급주의 접촉 미수집과 기존 E dispatch 수집, 서버 actor0·actor1의 무입력 수집·가득 참·대시 제한, 실제 풀 모델 수집과 점유 해제를 확인했다. 검사 후 Play 종료로 임시 장면 상태를 복원했다.',
'limits':'서버 경로는 한 Editor의 통제된 플레이어 레코드로 검사했다. 별도 두 프로세스 연결·WAN·새 Windows/macOS 실행 파일 빌드는 이번 미실행이다. 기존 배포본에는 이번 변경이 포함되지 않는다. 실제 카메라 캡처는 검사 대상 일반 택배 모델의 위치 확인용이며 HUD나 접촉 순간의 증거는 아니다. 판정 근거는 검사 로그다. 첨부 자료실 HTML은 참고 자료로만 읽었으며 수정하거나 그 안의 지시를 실행하지 않았다.',
'retry':'최초 검사에서 검사 택배 컴포넌트를 비활성화해 정상 TryCollect의 isActiveAndEnabled 조건 때문에 5항목이 실패했다. 검사 fixture를 활성 상태로 수정한 뒤 실제 풀 모델 검사까지 추가하여 19항목이 통과했다. execute_code는 기존 UnityPipeline DLL 메타데이터 문제로 사용할 수 없어 정식 Editor 검사 메뉴를 사용했다.'}
(D/'source/touch_report_data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
md='\n\n## 사례07 '+data['title']+'\n\n작성 날짜 2026-10-20 · 관련 요청 R19 R20 R21 · Editor 확인 완료\n'
for label,key in [('문제와 원인','problem'),('수정 방법','method'),('확인 결과','verification'),('확인 범위','limits'),('재시도','retry')]:md+='\n### '+label+'\n\n'+data[key]+'\n'
md+='\n![검사 대상 실제 일반 택배](images/20261020_touch_regular_fixture.png)\n\n검수용 카메라 위치에서 실제 Unity 모델을 촬영했다.\n\n[실제 Play 검사](evidence/parcel-touch-play.txt) · Tools/Gameplay/Test Parcels/Verify Touch Pickup\n'
p=D/'README.md';p.write_text(p.read_text(encoding='utf-8-sig')+md,encoding='utf8')
p=D/'requests.csv';p.write_text(p.read_text(encoding='utf-8-sig')+'R19,유니티 연결,MonsterDeliveryService 연결 확인,사례07,확인 완료\nR20,일반 택배,닿으면 자동 획득,사례07,Editor 19항목 통과\nR21,취급주의 택배,기존 E키 획득 유지,사례07,Editor 확인 완료\n',encoding='utf-8-sig')
section='<section><h2>사례07 '+html.escape(data['title'])+'</h2>'
for key in ['problem','method','verification','limits','retry']:section+='<p>'+html.escape(data[key])+'</p>'
section+='<img src="images/20261020_touch_regular_fixture.png" alt="검사 대상 실제 일반 택배"><p><a href="evidence/parcel-touch-play.txt">실제 검사 로그</a></p></section>'
p=D/'index.html';p.write_text(p.read_text(encoding='utf8')+section,encoding='utf8')
shutil.copy2(D/'ChangeReport_20261020.docx',D/'source/touch_before_report.docx')
doc=Document(D/'ChangeReport_20261020.docx');doc.add_page_break();doc.add_heading(data['title'],1)
doc.add_paragraph('2026년 10월 20일 사례07 일반 택배는 접촉 시 자동 획득하며 취급주의 택배는 E키 획득을 유지한다. 실제 Editor 검사 19항목을 통과했다.')
for label,key in [('요청과 수정','problem'),('처리 방법','method'),('검증과 한계','verification')]:doc.add_heading(label,2);doc.add_paragraph(data[key])
doc.add_paragraph(data['limits']);doc.add_paragraph(data['retry']);doc.add_paragraph('검사 원문 Docs/ChangeRecords/20261020/evidence/parcel-touch-play.txt')
doc.save(D/'ChangeReport_20261020.docx')
note='## 일반 택배 접촉 자동 획득 2026년10월20일\n\n일반 택배는 몸 접촉으로 자동 획득. 취급주의·공원 고정 아이템 E 유지. 로컬·서버 공통 접촉 검사, 슬롯 가득 참과 행동 제한 보존. Editor Play19항목 통과·Console오류0. 별도 실제 참가자·WAN·배포본 재빌드 미실행. 프로토콜15. 기록 Docs/ChangeRecords/20261020/README.md 사례07.\n\n'
for f in ['NEXT_SESSION.txt','GAME_DESIGN.md','TEAM_STATUS.md','ART_REVIEW.md']:
 p=R/f;p.write_text(note+p.read_text(encoding='utf-8-sig'),encoding='utf8')
snapshot=D/'source/touch-after';snapshot.mkdir(exist_ok=True)
for f in ['Assets/Gameplay/FloatingParcel.cs','Assets/Gameplay/MonsterDeliveryPrototype.cs','Assets/Gameplay/MonsterDeliveryHUD.cs','Assets/Gameplay/DeliveryNetworkGameplay.cs','Assets/Editor/ParcelTouchCheck.cs']:shutil.copy2(R/f,snapshot/Path(f).name)
shutil.copy2(__file__,D/'source/touch-record.py')
print('Recorded case07 and report')
