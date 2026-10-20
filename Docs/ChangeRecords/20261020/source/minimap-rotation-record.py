from pathlib import Path
from docx import Document
import json,html,shutil
R=Path('C:/Users/xlfhf/Documents/Codex/2026-09-18/https-github-com-kangminsu123-monsterdeliveryservice-git/MonsterDeliveryService');D=R/'Docs/ChangeRecords/20261020'
data={'title':'방장 미니맵 화살표 회전 시 위치 이동 수정',
'problem':'멀티에서 내 위치 좌표가 계속 바뀐다는 신고였다. 사용자는 방장 화면에서 방향을 돌릴 때 화살표 위치도 움직인다고 확인했다. DeliveryMinimap은 실제 로컬 Player Transform을 계속 참조하고 있었으며, 발견한 원인은 서버 좌표 교환이 아닌 GUI 회전 중심이었다.',
'cause':'기존 GUIUtility.RotateAroundPivot은 회전 행렬을 기존 GUI 배율 행렬 앞에 곱한다. 이 프로젝트는1920x1080 논리 좌표에 화면 배율을 적용한 뒤 논리 좌표 중심을 그대로 전달하여 배율1이 아니면 회전 중심이 화면상 화살표 중심과 어긋났다. 또한 맵 가장자리에서 DrawMarker는 위치를 안쪽으로 제한하지만 회전 중심은 제한 전 좌표를 사용했다. Unity 공식 UnityCsReference Modules/IMGUI/GUIUtility.cs의 구현과 기존 코드를 대조했다.',
'method':'PlayerMarkerMatrix에서 배율 행렬 뒤에 중심 이동·회전·복귀 행렬을 곱해 논리 좌표에서 화살표만 회전한다. MarkerCenter를 재사용해 그리기와 회전 모두 같은 가장자리 제한 좌표를 사용한다. 실제 플레이어 위치와 네트워크 이동·스냅샷·프로토콜15에는 변경이 없다.',
'verification':'Unity6000.5.10f1 실제 Editor Play의 MinimapRotationCheck로 생산 코드 PlayerMarkerMatrix를 호출했다. 배율0.5·2/3·1·4/3·2의5종과 중앙·좌상단·우하단3곳, 회전0·45·90·180·270·359도6종을 검사했다. 15개 조합별 판정(총90회 회전) 모두 통과했다. 고정 좌표의 중심 오차는 최대0.000252화면픽셀 미만이었다. 기존 Unity 회전 행렬 순서와 수정 전 가장자리 중심을 수학적으로 재현한 기준값도 로그에 함께 남겼다. 최종 Console오류0. 검사 후 Play종료.',
'limits':'별도 연결된 방장·참가자 두 프로세스와 사용자 해상도에서의 실플레이 재검증은 미실행이다. 이번 검사는 Editor 내 실제 생산 코드의 행렬 계산 검증이며 실제 화면 캡처나 네트워크 전송 검사가 아니다. 처음 OnGUI 기반 검사 helper는 GameView Repaint를 받지 못해 로그가 남지 않아 즉시 메뉴에서 생산 계산 코드를 호출하는 방식으로 바꿨다. 배포 실행파일 미재빌드. Word 본문은 갱신했으나 기존 LibreOffice 미설치로 페이지 화면검수 미완료.'}
(D/'source/minimap_rotation_report_data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
md='\n\n## 사례10 '+data['title']+'\n\n2026-10-20 · 요청R26 · Editor 확인 완료\n\n'+'\n\n'.join(data[k] for k in ['problem','cause','method','verification','limits'])+'\n\n[15조합90회전 검사 원문](evidence/minimap-rotation-after.txt)\n\n공식 구현 https://github.com/Unity-Technologies/UnityCsReference/blob/master/Modules/IMGUI/GUIUtility.cs\n'
p=D/'README.md';p.write_text(p.read_text(encoding='utf8')+md,encoding='utf8')
p=D/'requests.csv';p.write_text(p.read_text(encoding='utf-8-sig')+'R26,멀티 미니맵,방장이 방향을 돌릴 때 화살표 위치가 움직임,사례10,Editor 15조합90회전 통과\n',encoding='utf-8-sig')
p=D/'index.html';p.write_text(p.read_text(encoding='utf8')+'<section><h2>'+data['title']+'</h2>'+''.join('<p>'+html.escape(data[k])+'</p>' for k in ['problem','cause','method','verification','limits'])+'<p><a href="evidence/minimap-rotation-after.txt">검사 원문</a></p></section>',encoding='utf8')
doc=Document(D/'ChangeReport_20261020.docx');doc.add_page_break();doc.add_heading(data['title'],1)
for k in ['problem','cause','method','verification','limits']:doc.add_paragraph(data[k])
doc.save(D/'ChangeReport_20261020.docx')
note='## 방장 미니맵 회전 시 위치이동 수정 2026년10월20일\n\n해상도 배율과 회전중심 불일치 및 맵가장자리 중심 불일치 수정. 실제 좌표/네트워크 변경 없음. Editor 5배율×3위치×6방향90회전,15조합통과·Console0. Docs/ChangeRecords/20261020/README.md 사례10. 별도 연결된 방장/참가자 실플레이와 배포본 재빌드 미실행. Word 화면검수는 기존LibreOffice 미설치로 미완료.\n\n'
for f in ['NEXT_SESSION.txt','GAME_DESIGN.md','TEAM_STATUS.md','ART_REVIEW.md']:
 p=R/f;p.write_text(note+p.read_text(encoding='utf8'),encoding='utf8')
snap=D/'source/minimap-rotation-after';snap.mkdir(exist_ok=True)
for f in ['Assets/Gameplay/DeliveryMinimap.cs','Assets/Editor/MinimapRotationCheck.cs']:shutil.copy2(R/f,snap/Path(f).name)
shutil.copy2(__file__,D/'source/minimap-rotation-record.py')
print('case10 recorded')
