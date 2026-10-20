from docx.shared import Inches,Pt

def append_case04(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 for label,key in [('요청과 원인','cause'),('적용 방법','method'),('획득과 조작','pickup')]:
  p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(c['retry_short'] if key=='retry' else c[key])
 t=doc.add_table(rows=1,cols=2);t.style='Table Grid';t.autofit=False
 t.columns[0].width=Inches(1.2);t.columns[1].width=Inches(5.1)
 for cell,tx in zip(t.rows[0].cells,['신규 사용 아이템','효과와 처리']):cell.text=tx
 for name,effect,detail in c['effects']:
  cells=t.add_row().cells;cells[0].text=name;cells[1].text=effect+'. '+detail
 for row in t.rows:
  row.cells[0].width=Inches(1.2);row.cells[1].width=Inches(5.1)
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3)
    for r in p.runs:r.font.size=Pt(9)
 for label,key in [('멀티 동기화','network'),('검증 결과','verification'),('재시도','retry'),('확인 범위','limits')]:
  p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(c['retry_short'] if key=='retry' else c[key])
 doc.add_page_break();doc.add_heading('공원 일렬 배치와 실제 상대 효과',1)
 doc.add_paragraph('2026년10월19일 실제 Unity 화면. 검수용으로 플레이어를 공원에 옮겼다. 장비3슬롯은 그대로 두고 촬영 때 손 모델만 잠시 숨겨5종이 모두 보이게 했다.')
 doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_special_five_item_row.png'),width=Inches(6.35))
 doc.add_paragraph('5종 고정 아이템을 일렬 배치했다. 각 모델을 바라보고 E로 장착한 뒤 좌클릭으로 사용한다.','Caption')
 doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_special_lifejacket_remote.png'),width=Inches(6.35))
 doc.add_paragraph('별도 참가자가 구명조끼를 사용한 뒤 호스트 카메라에서 본 모습. 면역은5초 동안 유지되고, 종료 후 방송기 기절이 다시 적용되는 것을 로그로 확인했다. evidence/special-row-connected/host.txt 및 client.txt.','Caption')
