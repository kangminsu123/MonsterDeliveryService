from docx.shared import Inches
def append_case05(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 for label,key in [('요청','problem'),('원인','cause'),('수정 방법','method'),('판정 정책','policy'),('멀티 동기화','network'),('검증 결과','verification'),('재시도','retry'),('확인 범위','limits')]:
  p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(c[key])
 doc.add_page_break();doc.add_heading('실제 바닥과 벽 충돌 폭발 화면',1)
 doc.add_paragraph('2026년10월19일 실제 Unity 캡처다. 검수용 위치와 카메라를 설정했다. 정지 화면은 연출 확인용이며 피격 범위·상승·동기화·시간은 실행 로그로 검증했다.')
 doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_air_cannon_burst_floor.png'),width=Inches(6.35))
 doc.add_paragraph('바닥 적중 후 공기 파편과3개의 충격파가 퍼지는 순간이다. 파동은0.3초에 반경6m까지 확장되며 이 화면은 확장 중이다. 이펙트에는 충돌체가 없다.','Caption')
 doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_air_cannon_burst_client_wall.png'),width=Inches(6.35))
 doc.add_paragraph('별도 참가자가 받은 벽 적중 폭발. 서버가 보낸 충돌 위치와 표면 법선을 사용해 벽 방향의 파동과 공기 파편을 재생한다. evidence/air-cannon-burst-connected/host.txt와client.txt 참조.','Caption')
