from docx.shared import Inches

def append_case06(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 for label,key in [('요청','problem'),('원인','cause'),('수정 방법','method'),('판정 정책','policy'),('멀티 동기화','network'),('검증 결과','verification'),('재시도','retry'),('확인 범위','limits')]:
  p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(c[key])
 doc.add_paragraph('검증 원문 evidence/broadcaster-offline.txt · broadcaster-special-regression.txt · broadcaster-row-regression.txt · broadcaster-connected/host.txt 및 client.txt')
 doc.add_page_break();doc.add_heading('실제 방송기 사용과 참가자 화면',1)
 doc.add_paragraph('2026년 10월 19일 실제 연결된 두 Unity Editor에서 촬영했다. 검수용 위치와 카메라를 사용했다. 원격 참가자 화면에는 동일한 방송 동작의 경과 시간을 적용한다.')
 for filename,caption in [('20261019_broadcaster_host_connected.png','사용자 시점. 여러 겹의 따뜻한 음파와 금빛 불꽃이 퍼지고, 맞은 상대 머리 위에 기절 표시가 나타난다. 연출은 1.3초, 기절은 4초로 서로 구분한다.'),('20261019_broadcaster_client_connected.png','별도 참가자의 검수용 외부 카메라. 수평 음파 3겹·수직 반원 2개·중앙 경고 원·방사형 음파가 보인다. 파동은 최대 12m까지 확장한다. 광원과 불꽃의 작동·소멸은 실행 로그로도 확인했다.')]:
  doc.add_paragraph().add_run().add_picture(str(D/'images'/filename),width=Inches(6.35))
  doc.add_paragraph(caption,'Caption')
