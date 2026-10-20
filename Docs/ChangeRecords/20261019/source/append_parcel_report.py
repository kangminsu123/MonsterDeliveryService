from docx.shared import Inches
def append_case02(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 for title,key in [('문제와 요청','problem'),('제작 방법','method')]:
  p=doc.add_paragraph();p.add_run(title+'  ').bold=True;p.add_run(c[key])
 doc.add_paragraph().add_run().add_picture(str(D/'images'/c['image']),width=Inches(4.1))
 doc.add_paragraph('실제 Blender 상세 카메라 출력. 글자 없이 잔 모양 파손주의·위쪽 화살표·경고 띠만 남겼다.','Caption')
 for title,key in [('결과와 검증','result'),('적용 범위와 저장','limits')]:
  p=doc.add_paragraph();p.add_run(title+'  ').bold=True;p.add_run(c[key])
 doc.add_paragraph('최종 검증 evidence/parcel_marks_audit.json · 초기 검증 fragile_final_audit.json · 편집 원문 source/case02/')
