from docx.shared import Inches,Pt

def append_case03(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 for label,key in [('문제와 원인','cause'),('적용 방법','method'),('공원 택배','parcel'),('시작 장비','loadout')]:
  p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(c[key])
 t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
 for cell,tx in zip(t.rows[0].cells,['테스트 아이템','대상과 범위','테스트 효과']):cell.text=tx
 for row in c['items']:
  for cell,tx in zip(t.add_row().cells,row):cell.text=tx
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3)
    for r in p.runs:r.font.size=Pt(9)
 for label,key in [('검증 결과','verification'),('범위와 제한','limits'),('재시도','retry')]:
  p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(c[key])
 doc.add_page_break();doc.add_heading('유니티 실제 적용 화면',1)
 doc.add_paragraph('2026년10월19일 SampleScene 실제 실행 화면이다. 촬영을 위해 플레이어를 공원으로 옮겼으며 게임의 네 모서리 랜덤 시작 위치는 유지한다.')
 doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_special_park_test.png'),width=Inches(6.35))
 doc.add_paragraph('본인 시점. 공원 택배1개와 대포·방송기·포장기 각1회가 보인다. 슬롯이 가득 찬 상태라 G로 비운 뒤 E로 열 수 있다.','Caption')
 doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_special_connected_packing_client.png'),width=Inches(6.35))
 doc.add_paragraph('별도 참가자 Editor의 실제 카메라. 통제된 공중 시험 지면에서 상대에게 포장기 사용 후, 마크가 있는 상자와 기절 표시가 함께 보인다. 정지 화면의 시간 검증은 evidence/special-connected/*.txt와 csv에 따로 남겼다.','Caption')
