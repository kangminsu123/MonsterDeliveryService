from docx.shared import Inches
def append_case05(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 for label,key in [('요청','problem'),('원인','cause'),('수정 방법','method'),('판정 정책','policy'),('멀티 동기화','network'),('검증 결과','verification'),('재시도','retry'),('확인 범위','limits')]:
  p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(c[key])
 doc.add_page_break();doc.add_heading('실제 비행과 착지 범위 폭발 화면',1)
 doc.add_paragraph('2026년10월19일 실제 Unity 캡처다. 검수용으로 위치와 카메라를 설정했다. 정지 화면은 표시 확인용이며 효과와 시간은 실행 로그로 검증했다.')
 doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_air_cannon_flight.png'),width=Inches(6.35))
 doc.add_paragraph('본인 발사 후 중력으로 날아가는 공기 탄. 발사 즉시 범위 피격은 발생하지 않는다. 멀리 있는 보라색 상대는 착지 범위 밖이다.','Caption')
 doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_air_cannon_client_impact.png'),width=Inches(6.35))
 doc.add_paragraph('별도 참가자가 수신한 착지 원형 파동. 서버에서 호스트가 발사한 탄의 충돌과5m 범위를 확정하고 참가자에게 전달했다. evidence/air-cannon-connected/host.txt와client.txt 참조.','Caption')
