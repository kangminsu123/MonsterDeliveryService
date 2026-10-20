from pathlib import Path
import json,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn
W=Path('C:\\Users\\xlfhf\\Documents\\Codex\\2026-10-17\\dbslx');R=Path('C:\\Users\\xlfhf\\Documents\\Codex\\2026-09-18\\https-github-com-kangminsu123-monsterdeliveryservice-git\\MonsterDeliveryService');D=Path(__file__).resolve().parent.parent
data=json.loads((D/"source/report_data.json").read_text(encoding="utf-8"))
doc=Document(R/'Docs/ChangeRecords/20261018/source/report_template.docx')
for e in list(doc._element.body):
 if e.tag!=qn('w:sectPr'):doc._element.body.remove(e)
for sec in doc.sections:
 sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
 for area in [sec.header,sec.footer]:
  for p in area.paragraphs:p.text=''
for n in ['Normal','Title','Heading 1','Caption']:
 st=doc.styles[n];st.font.name='맑은 고딕';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
for e in list(doc.styles.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
doc.styles['Normal'].font.size=Pt(9.5);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.08;doc.styles['Title'].font.size=Pt(22);doc.styles['Heading 1'].font.size=Pt(14)
doc.add_paragraph(data['title'],'Title');doc.add_paragraph(data['date']);doc.add_paragraph(data['intro'])
for label,key in [('문제와 요청','problem'),('제작 방법','method')]:
 p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(data[key])
t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
for c,tx in zip(t.rows[0].cells,['아이템','시각적 특징','삼각형']):c.text=tx
for name,desc,count in data['items']:
 for c,tx in zip(t.add_row().cells,[name,desc,str(count)]):c.text=tx
for row in t.rows:
 for c in row.cells:
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(4)
   for run in p.runs:run.font.size=Pt(9)
for label,key in [('검증 결과','verification'),('적용 범위','limits')]:
 p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(data[key])
doc.add_page_break();doc.add_heading('실제 블렌더 디자인 화면',1)
doc.add_paragraph('2026년10월19일 연결된 Blender5.1에서 제작하고 직접 렌더한 비교 화면이다. 검수용 받침대와 조명은 모델과 분리돼 있으며, 그림은 인게임 캡처가 아니다.')
doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_special_items_blender_design.png'),width=Inches(6.35))
doc.add_paragraph('위쪽 대포·방송기·포장기, 아래쪽 구명조끼·배터리.5종은 각각 독립 루트로 편집할 수 있다.','Caption')
for label,key in [('저장 상태','save'),('재시도','retry')]:
 p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(data[key])
doc.add_paragraph('검증 원문 evidence/final_audit.json · 편집 설계 source/special_items_passport.txt · 생성·모델링 원문 source/*.py')
from append_parcel_report import append_case02
append_case02(doc,D,data['case02'])
from append_import_report import append_case03
append_case03(doc,D,data['case03'])
from append_row_report import append_case04
append_case04(doc,D,data['case04'])
doc.save(D/'ChangeReport_20261019.docx');shutil.copy2(D/'ChangeReport_20261019.docx',W/'outputs/SpecialItemDesignReport.docx')
