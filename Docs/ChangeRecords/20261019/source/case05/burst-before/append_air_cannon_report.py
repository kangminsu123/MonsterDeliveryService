from docx.shared import Inches
def append_case05(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 for label,key in [('요청','problem'),('원인','cause'),('수정 방법','method'),('판정 정책','policy'),('멀티 동기화','network'),('검증 결과','verification'),('재시도','retry'),('확인 범위','limits')]:
  p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(c[key])
 doc.add_page_break();doc.add_heading('실제 직선 비행과 본인 발밑 발사 화면',1)
 doc.add_paragraph('2026년10월19일 실제 Unity 캡처다. 검수용 위치와 카메라를 설정했다. 정지 화면은 표시 확인용이며 직선 궤적·피격·상승·시간은 실행 로그로 검증했다.')
 doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_air_cannon_linear_flight.png'),width=Inches(6.35))
 doc.add_paragraph('바닥을 조준해 빠르게 직선으로 날아가는 공기 탄. 중력이나 추가 상승 속도를 사용하지 않는다. 닿은 위치의 반경5m에서 폭발한다.','Caption')
 doc.add_paragraph().add_run().add_picture(str(D/'images/20261019_air_cannon_linear_client_self_jump.png'),width=Inches(6.35))
 doc.add_paragraph('별도 참가자가 자기 발밑 바닥에 발사한 뒤 위로 날아가는 실제 순간. 서버가 참가자 본인의 힘을 확정하고 양쪽에서 상승 위치를 확인했다. evidence/air-cannon-linear-connected/host.txt와client.txt 참조.','Caption')
