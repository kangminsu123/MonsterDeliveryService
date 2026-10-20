from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
import html,re

def append_case02(doc,D,data):
    if 'Case2Body' not in doc.styles:
        style=doc.styles.add_style('Case2Body',WD_STYLE_TYPE.PARAGRAPH)
        style.base_style=doc.styles['Normal'];style.font.size=Pt(10.5)
        style.paragraph_format.space_after=Pt(6);style.paragraph_format.line_spacing=1.05
    def p(text):return doc.add_paragraph(text,'Case2Body')
    def image(name,caption,width=6.6):
        q=doc.add_paragraph();q.add_run().add_picture(str(D/'images'/name),width=Inches(width))
        q.paragraph_format.space_after=Pt(2);q.alignment=1
        doc.add_paragraph(caption,'Caption')
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        widths=[2.7,1.8,2.1] if len(headers)==3 else [2.3,4.3]
        for c,text in zip(t.rows[0].cells,headers):c.text=text
        for values in rows:
            for c,text in zip(t.add_row().cells,values):c.text=str(text)
        for ri,row in enumerate(t.rows):
            for ci,cell in enumerate(row.cells):
                cell.width=Inches(widths[ci]);tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                tcpr.append(borders)
                if ri==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EAF0');tcpr.append(shade)
                for q in cell.paragraphs:
                    q.paragraph_format.space_after=Pt(5);q.paragraph_format.space_before=Pt(5)
                    for r in q.runs:r.font.size=Pt(10);r.font.bold=ri==0;r.font.color.rgb=RGBColor(0,0,0)
    doc.add_page_break()
    doc.add_paragraph('참가자 지연과 개인전 UI 수정','Title')
    p('2026년 10월 16일  |  사례 02  |  요청 R05부터 R09')
    p(data['intro'])
    p('기존 노란 고정 PNG가 배정색을 읽지 않았고, 팀 문구도 1대1을 구분하지 않았다.')
    image(data['images'][0][0],data['images'][0][1])
    table(['측정 항목','수정 전','수정 후'],data['metrics'])
    p(data['measurement_scope'])

    doc.add_page_break();doc.add_heading('원인 확인과 수정 방법',1)
    for text in data['causes'][1:]:p(text)
    image('20261016_latency_movement_trace.png','실제 프레임별 CSV로 그린 이동 거리. 빨강은 서버 응답 대기, 초록은 로컬 예측이다. 네트워크 지연을 없앴다는 뜻은 아니다.',6.0)
    for text in data['methods']:p(text)

    doc.add_page_break();doc.add_heading('수정 후 개인전 화면',1)
    image(data['images'][1][0],data['images'][1][1])
    p(data['ui_result'])
    doc.add_heading('메시지 누락과 위치 확인',2);p(data['loss'])
    table(['검사','최종 결과'],data['checks'][:2])
    doc.add_heading('재시도와 수정 이력',2);p(data['retries'])

    doc.add_page_break();doc.add_heading('팀 경기 회귀와 새 실행 파일',1)
    image(data['images'][2][0],data['images'][2][1])
    table(['실행 검사','최종 결과'],data['checks'][2:])
    p(data['build']);p(data['delivery']);p(data['limitations'])
    render_web(D,data)

def render_web(D,data):
    marker='<!-- CASE02_UI_NETWORK -->'
    readme=D/'README.md';base=readme.read_text(encoding='utf-8').split(marker)[0].rstrip()
    base=base.replace('2026-10-16 · 사례 01 · R01~R04','2026-10-16 · 사례 01·02 · R01~R09')
    lines=[marker,'## 사례 02 참가자 지연과 개인전 UI 수정','2026-10-16 · R05~R09 · 결함 수정 및 요구 변경 · 검증 완료',data['intro'],'### 문제와 원인',*data['causes'],'### 확인 과정과 측정 범위',data['measurement_scope']]
    lines+=['| 측정 항목 | 수정 전 | 수정 후 |','|---|---|---|']+['| '+' | '.join(row)+' |' for row in data['metrics']]
    lines+=['### 수정 방법',*data['methods'],'### 실제 화면']
    for name,cap in data['images']+[('20261016_latency_movement_trace.png','프레임별 실제 CSV의 이동 거리 비교'),('20261016_latency_client_after_loss.png','송신 앱 메시지20% 누락을 추가한 실제 참가자 검수 화면')]:
        lines+=['!['+cap+'](images/'+name+')',cap]
    lines+=['### 수정 후 검증',data['ui_result'],data['loss']]
    lines+=['- '+a+': '+b for a,b in data['checks']]
    lines+=['### 재시도',data['retries'],'### 빌드와 적용',data['build'],data['delivery'],'### 남은 제한',data['limitations'],'### 근거와 소스']
    lines+=['- ['+label+']('+path+')' for label,path in data['evidence']]
    lines+=['- `'+name+'`' for name in data['files']]
    text=base+'\n\n'+'\n\n'.join(lines)+'\n'
    text=re.sub(r'(?m)^(\|[^\n]+\|)\n\n(?=\|)',r'\1\n',text)
    readme.write_text(text,encoding='utf-8')
    web=D/'index.html';base=web.read_text(encoding='utf-8').split(marker)[0]
    base=base.replace('</body>','').replace('</html>','')
    def para(s):return '<p>'+html.escape(s)+'</p>'
    blocks=[marker,'<article id="case02"><h1>사례 02 참가자 지연과 개인전 UI 수정</h1>',para(data['intro'])]
    for title,texts in [('문제와 원인',data['causes']),('확인 과정', [data['measurement_scope']]),('수정 방법',data['methods'])]:
        blocks+=['<h2>'+title+'</h2>']+[para(t) for t in texts]
    blocks+=['<table style="border-collapse:collapse"><tr><th>측정 항목</th><th>수정 전</th><th>수정 후</th></tr>']
    blocks+=['<tr>'+''.join('<td style="border:1px solid #d9d9d9;padding:10px">'+html.escape(c)+'</td>' for c in row)+'</tr>' for row in data['metrics']]
    blocks+=['</table>']
    for name,cap in data['images']+[('20261016_latency_movement_trace.png','실제 CSV의 이동 거리 비교'),('20261016_latency_client_after_loss.png','메시지20% 누락 검수용 재현')]:
        blocks+=['<figure><img style="max-width:100%" src="images/'+name+'"><figcaption>'+html.escape(cap)+'</figcaption></figure>']
    for title,texts in [('결과',[data['ui_result'],data['loss']]),('검증',[a+': '+b for a,b in data['checks']]),('재시도',[data['retries']]),('배포와 제한',[data['build'],data['delivery'],data['limitations']])]:
        blocks+=['<h2>'+title+'</h2>']+[para(t) for t in texts]
    blocks+=['<ul>']+['<li><a href="'+path+'">'+html.escape(label)+'</a></li>' for label,path in data['evidence']]+['</ul></article></body></html>']
    web.write_text(base+'\n'.join(blocks),encoding='utf-8')
