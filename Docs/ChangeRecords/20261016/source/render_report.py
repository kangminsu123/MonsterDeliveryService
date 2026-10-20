import sys
sys.dont_write_bytecode=True
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml.ns import qn
D=Path(__file__).resolve().parent.parent
data=json.loads((D/'source/report_data.json').read_text(encoding='utf-8'))
globals().update(data)
doc=Document(D/'source/report_template.docx')
for child in list(doc._element.body):
 if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
def p(t):doc.add_paragraph(t)
def pic(name,caption,width=6.6):
 q=doc.add_paragraph();q.add_run().add_picture(str(D/'images'/name),width=Inches(width));q.paragraph_format.space_after=Pt(2)
 doc.add_paragraph('2026-10-16 실제 Unity 검수용 캡처 · '+caption,'Caption')
def table(rows):
 t=doc.add_table(rows=0,cols=2);t.autofit=False
 for a,b in rows:
  cells=t.add_row().cells
  for c,text,w in zip(cells,[a,b],[1.25,5.35]):
   c.text=text;c.width=Inches(w)
   for q in c.paragraphs:
    q.paragraph_format.space_after=Pt(4)
    for r in q.runs:r.font.size=Pt(9.5)
doc.add_paragraph('택배 배치와 멀티플레이 수정 보고서' if 'case02' in data else '택배 무작위 배치와 재생성 변경 보고서','Title');p('2026년 10월 16일  |  Monster Delivery Service  |  사례 01')
p(intro);doc.add_heading('요청과 변경 배경',1);p(problem);table(params);pic(*images[0]);p('청록은 택배가 있는 후보, 노랑은 빈 후보다. 후보 번호와 상태 색은 검수용이며 일반 플레이에는 표시하지 않는다.')
doc.add_page_break();doc.add_heading('획득 후 배치와 실제 택배',1);p('택배를 먹은 위치는 비우고 다른 빈 후보에 보충한다. 아래 상공 화면의 빨강은 획득 위치, 초록은 새 위치다.');pic(*images[1]);p(layout)
t=doc.add_table(rows=1,cols=2);t.autofit=False
for cell,(name,cap) in zip(t.rows[0].cells,images[2:]):
 cell.width=Inches(3.3);q=cell.paragraphs[0];q.add_run().add_picture(str(D/'images'/name),width=Inches(3.2));cell.add_paragraph('2026-10-16 · '+cap,'Caption')
p('기존 상자·보호막·부유·흰 테두리·E 획득을 재사용했다. 원본 크기와 아이템 9종 동일 확률 지급은 유지한다.')
doc.add_page_break();doc.add_heading('검증 결과와 적용 방법',1);table([(a,b) for a,b,_ in tests]);p(verification);doc.add_heading('구현과 동기화',2);p(method);p(network);doc.add_heading('재시도와 빌드 환경',2);p(retry);p(build);doc.add_heading('새 실행 파일 적용',2);p('MonsterDeliveryService_20261016.zip 전체를 압축 해제한 뒤 MonsterDeliveryService.exe를 실행한다. 1대1과 2대2 메뉴 및 Tailscale 주소 입력 방식은 유지한다. 택배 동기화 형식이 바뀌었으므로 모든 참가자는 이번 새 버전을 사용한다. 원본 소스·후보 좌표·캡처·전체 PASS 로그는 날짜별 기록 폴더에 보존했다.')
if 'case02' in data:
 from append_ui_network_report import append_case02
 append_case02(doc,D,data['case02'])
if 'case03' in data:
 from append_mac_report import append_case03
 append_case03(doc,D,data['case03'])
for q in doc.paragraphs:
 q.paragraph_format.keep_together=True
doc.styles['Normal'].font.size=Pt(9.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Normal'].paragraph_format.line_spacing=1
doc.styles['Title'].font.size=Pt(22)
for q in doc.paragraphs:
 if q._p.xpath('.//w:br[@w:type="page"]'):
  q.paragraph_format.keep_together=False
  q.paragraph_format.space_after=Pt(0);q.paragraph_format.space_before=Pt(0);q.paragraph_format.line_spacing=Pt(1)
  for r in q.runs:r.font.size=Pt(1)
doc.save(D/'ChangeReport_20261016.docx')
