import html
from docx.shared import Pt

MARKER='<!-- CASE03_MACOS -->'

def append_case03(doc,D,data):
    doc.add_page_break()
    doc.add_paragraph('맥 개발 테스트 패키지','Title')
    def p(t):return doc.add_paragraph(t,'Case2Body')
    p('2026년 10월 16일  |  사례 03  |  요청 R10')
    p(data['intro'])
    for title,key in [('빌드 환경과 처리','method'),('실제 파일 검증','verification'),('실행과 연결','launch'),('재시도와 진단','retry'),('남은 확인','limits')]:
        doc.add_heading(title,2);p(data[key])
    p('근거: evidence/build-macos.json, evidence/package-verification-macos.json, evidence/macos-build-diagnostics.txt. 실행 안내와 BUILD_INFO.json은 맥 ZIP에 포함했다.')
    readme=D/'README.md';base=readme.read_text(encoding='utf-8').split(MARKER)[0].rstrip()
    parts=[MARKER,'## 사례 03 맥 개발 테스트 패키지','2026-10-16 · R10 · 요구 변경 · 패키징 완료 및 실기 미검증',data['intro']]
    for title,key in [('원인과 방법','method'),('검증 결과','verification'),('실행과 연결','launch'),('재시도','retry'),('남은 제한','limits')]:parts+=['### '+title,data[key]]
    parts+=['### 근거','- [빌드 결과](evidence/build-macos.json)','- [앱 구조와 ZIP 검사](evidence/package-verification-macos.json)','- [진단과 재시도](evidence/macos-build-diagnostics.txt)','- [재생성 원문](source/report_data.json)']
    readme.write_text(base+'\n\n'+'\n\n'.join(parts)+'\n',encoding='utf-8')
    web=D/'index.html';base=web.read_text(encoding='utf-8').split(MARKER)[0].replace('</body>','').replace('</html>','').rstrip()
    blocks=[MARKER,'<article id="case03"><h1>사례 03 맥 개발 테스트 패키지</h1>','<p>'+html.escape(data['intro'])+'</p>']
    for title,key in [('원인과 방법','method'),('검증 결과','verification'),('실행과 연결','launch'),('재시도','retry'),('남은 제한','limits')]:blocks+=['<h2>'+title+'</h2><p>'+html.escape(data[key])+'</p>']
    blocks+=['<p><a href="evidence/build-macos.json">빌드 결과</a> · <a href="evidence/package-verification-macos.json">앱 구조와 ZIP 검사</a> · <a href="evidence/macos-build-diagnostics.txt">진단과 재시도</a></p></article></body></html>']
    web.write_text(base+'\n'+'\n'.join(blocks),encoding='utf-8')
