from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt
D=Path(__file__).resolve().parent.parent
backup=D/'source/ChangeReport_before_network.docx'
n=json.loads((D/'source/report_data.json').read_text(encoding='utf-8'))['network']
globals().update(n)
doc=Document(backup);doc.paragraphs[0].text='아이템과 멀티플레이 구현 보고서';doc.paragraphs[0].style='Title'
def p(t):doc.add_paragraph(t)
def pic(key,cap):
 q=doc.add_paragraph();q.add_run().add_picture(str(D/f'images/20261015_{key}.png'),width=Inches(6.6));q.paragraph_format.space_after=Pt(2)
 c=doc.add_paragraph('2026-10-15 실제 네트워크 검수 세션 · '+cap,'Caption');c.paragraph_format.space_after=Pt(6)
def table(rows):
 t=doc.add_table(rows=0,cols=2);t.autofit=False
 for a,b in rows:
  c=t.add_row().cells;c[0].text=a;c[1].text=b;c[0].width=Inches(1.2);c[1].width=Inches(5.4)
  for cell in c:
   for q in cell.paragraphs:
    q.paragraph_format.space_after=Pt(3)
    for r in q.runs:r.font.size=Pt(9.5)
doc.add_page_break();doc.add_heading('사례 02 멀티플레이와 실행 패키지',1);p(intro);table(params);pic(*images[0]);p('원인과 처리: '+problem)
doc.add_page_break();doc.add_heading('2대2 대기방과 결과',1);p('네 명이 입장하면 A/B 각 두 명으로 배정한다. 아래는 최종 코드의 실제 연결 검수 세션이다.');pic(*images[1]);pic(*images[2])
doc.add_page_break();doc.add_heading('1대1 대기방과 결과',1);p('1대1은 정원이 두 명이며 참가자마다 다른 옷색을 자동 배정한다.');pic(*images[3]);pic(*images[4])
doc.add_page_break();doc.add_heading('실제 연결 상태와 아이템 피격',1);p('동료/상대는 이름 위 A/B로 구분한다. 아이스크림 투척은 클라이언트 입력에서 시작해 호스트 물리 판정과 다른 창의 상태 표시로 이어진다.');pic(*images[5]);pic(*images[6])
doc.add_page_break();doc.add_heading('검증 결과와 테스트 실행 방법',1);p(verification)
table([['1대1 최종 EXE','호스트 13개와 클라이언트 3개 통과'],['2대2 최종 EXE','호스트 14개와 클라이언트 각 3개 통과'],['방 수명 주기','재경기 대기방·나가기·재개설·메뉴 4개 통과']]);doc.add_heading('실행 순서',2)
p('ZIP을 전부 압축 해제하고 MonsterDeliveryService.exe를 실행한다. 1대1 또는 2대2를 고른 뒤 테스트방 만들기를 누른다. 다른 창에서는 같은 PC라면 127.0.0.1, 같은 LAN의 다른 PC라면 방장이 복사한 MD 코드를 입력해 참가한다. 함께 제공한 Test_1v1.cmd와 Test_2v2.cmd는 해당 수의 창을 자동으로 연다. EXE 옆의 Data·MonoBleedingEdge·DLL 폴더와 파일은 함께 유지한다.')
p('조작은 WASD 이동, 마우스 시점·좌클릭 사용, Space 점프, C 웅크리기, Shift 대쉬, E 상호작용, 1~3 슬롯, G 폐기, Esc 메뉴, F2 개발 도구다. 현재 개발용 아이템은 무제한이다.')
doc.add_heading('구현과 환경',2);p(method);p(build);p(retry);doc.add_heading('검증 범위와 남은 제한',2);p(limits)
active=False
for q in doc.paragraphs:
 if q.text=='검증 결과와 테스트 실행 방법':active=True
 if active and q.style.name in ('Normal','Caption'):
  q.paragraph_format.space_after=Pt(4);q.paragraph_format.line_spacing=1
  for r in q.runs:r.font.size=Pt(9.5)
doc.save(D/'ChangeReport_20261015.docx')
