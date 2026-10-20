from docx.shared import Inches

def append_case05(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 def paragraph(title,text):
  p=doc.add_paragraph();p.add_run(title+'  ').bold=True;p.add_run(text);p.paragraph_format.keep_together=True
 for title,key in [('문제와 기대 동작','problem'),('원인','cause'),('수정 방법','method'),('해결 결과','result')]:paragraph(title,c[key])
 doc.add_heading('실제 Editor 화면',1)
 t=doc.add_table(rows=1,cols=2);t.autofit=False
 for cell,(name,caption) in zip(t.rows[0].cells,c['images']):
  cell.paragraphs[0].add_run().add_picture(str(D/'images'/name),width=Inches(3.05));cell.add_paragraph(caption,'Caption')
 for title,key in [('검증과 한계','verification'),('재시도와 근거','retry'),('변경 파일','files')]:paragraph(title,c[key])
