from docx.shared import Inches,Pt

def append_case02(doc,D,c):
 def p(text):
  paragraph=doc.add_paragraph(text);paragraph.paragraph_format.keep_together=True
 def section(title,key):doc.add_heading(title,1);p(c[key])
 def pic(name,caption):
  doc.add_paragraph().add_run().add_picture(str(D/'images'/name),width=Inches(6.0))
  doc.add_paragraph(caption,'Caption')
 doc.add_page_break();doc.add_heading(c['title'],1);p(c['date'])
 for title,key in [('문제와 기대 동작','problem'),('원인','cause'),('수정한 방법','method'),('수정 후 해결 결과','result')]:section(title,key)
 doc.add_page_break();doc.add_heading('본인 화면의 준비와 발사',1);p(c['capture_note'])
 for name,caption in c['images'][:2]:pic(name,caption)
 p('오른손 준비 0.16초 · 실제 발사 0.22초 · 복귀 0.55초 · 재사용 0.6초')
 doc.add_page_break();doc.add_heading('상대 화면의 오른손 투척',1);p('실제 참가자가 받은 호스트 슬라임 투척이다. 전경의 슬라임은 참가자 본인의 장비이며 가운데 노란 상대의 오른팔과 투사체가 검수 대상이다.')
 for name,caption in c['images'][2:]:pic(name,caption)
 doc.add_page_break()
 for title,key in [('투척 검증 결과와 재실행','verification'),('진단과 남은 확인','retry'),('적용 범위','limits'),('변경 파일','files')]:section(title,key)
 p('evidence/throw-motion-check.txt · 상태 재현98항목\nevidence/throw-connected-host.txt · 실제 호스트13항목\nevidence/throw-connected-client.txt · 실제 참가자13항목\nevidence/boxing-regression-after-throw.txt · 복싱 회귀25항목\nevidence/throw-session.diff · 이번 변경만 비교\nsource/case02 · 이전 보고서와 변경 전후 소스\nsource/report_data.json · 편집 원문')
