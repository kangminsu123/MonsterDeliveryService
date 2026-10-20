from docx.shared import Inches

def append_case03(doc,D,c):
 def p(text):
  paragraph=doc.add_paragraph(text);paragraph.paragraph_format.keep_together=True
 def section(title,key):doc.add_heading(title,1);p(c[key])
 def pic(name,caption,width=6):
  doc.add_paragraph().add_run().add_picture(str(D/'images'/name),width=Inches(width));doc.add_paragraph(caption,'Caption')
 doc.add_page_break();doc.add_heading(c['title'],1);p(c['date'])
 for title,key in [('문제와 기대 동작','problem'),('원인','cause'),('수정한 방법','method'),('수정 후 해결 결과','result')]:section(title,key)
 doc.add_page_break();doc.add_heading('실제 멀티의 비행 중 숨김과 복원',1);p(c['capture_note'])
 for name,caption in c['images'][:2]:pic(name,caption)
 doc.add_page_break();doc.add_heading('패킷 사이 비행 위치의 수정 전후',1);p(c['metrics_note']);pic(*c['images'][2])
 doc.add_heading('실측 비교',1)
 t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
 for cell,text in zip(t.rows[0].cells,['종류','수정 전 최대 위치 오차','수정 후 최대 위치 오차']):cell.text=text
 for row in [('뚫어뻥','2.31598m','0.03065m'),('슬라임','1.11920m','0.00003m'),('아이스크림','1.13438m','0.00003m'),('만년필','2.90329m','0.00000m')]:
  for cell,text in zip(t.add_row().cells,row):cell.text=text
 p('표는 상태 재현 입력의 기준 궤적에 대한 최대 오차이며 소수5자리 표시다. 실제 전송 지연이나 실제 충돌 오차를 뜻하지 않는다. 수정 전후 프레임 수가 달라 오차의 관측 최댓값에도 샘플링 영향이 있다.')
 p('실제 참가자 연결의 패킷 사이 이동/정지 프레임은 뚫어뻥145/2, 슬라임218/0, 아이스크림218/0, 만년필130/0이다. 원본 위치와 수신 순서는 evidence/projectile-connected-client.csv에 보존했다.')
 doc.add_page_break()
 for title,key in [('검증과 재실행','verification'),('재시도와 진단','retry'),('적용 범위와 제한','limits'),('변경 파일','files')]:section(title,key)
 p('evidence/projectile-display-before.txt · 이전 실패\nevidence/projectile-display-after.txt · 수정 후26항목\nevidence/projectile-connected-host.txt 및 client.txt · 실제 양쪽 결과\nevidence/throw-regression-after-projectile.txt · 기존 투척98항목 회귀\nevidence/projectile-session.diff · 이번 변경만 비교\nsource/case03 · 이전 보고서와 변경 전후 원문')
