from pathlib import Path
import base64, html, json, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

D=Path(__file__).resolve().parent.parent
O=Path(r'C:/Users/xlfhf/Documents/Codex/2026-10-17/dbslx/outputs')
c=json.loads((D/'source/item_usage_data.json').read_text(encoding='utf-8'))
doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 style=doc.styles[name];style.font.name='맑은 고딕';style.font.color.rgb=RGBColor(0,0,0)
 style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(23);doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 1'].paragraph_format.space_before=Pt(12);doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].paragraph_format.space_after=Pt(5)
doc.add_paragraph(c['title'],'Title');doc.add_paragraph(c['date']);doc.add_paragraph(c['intro'])
doc.add_heading('아이템별 획득당 사용량',1)
t=doc.add_table(rows=1,cols=4);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
widths=[1.2,.75,3.25,2.0]
for col,width in zip(t.columns,widths):col.width=Inches(width)
for cell,text in zip(t.rows[0].cells,['아이템','사용량','효과','소모 시점']):cell.text=text
repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
for row in c['rows']:
 for cell,text in zip(t.add_row().cells,row):cell.text=text
borders=OxmlElement('w:tblBorders')
for edge in ['top','left','bottom','right','insideH','insideV']:
 e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
t._tbl.tblPr.append(borders)
for i,row in enumerate(t.rows):
 for j,cell in enumerate(row.cells):
  cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  props=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');m.append(e)
  props.append(m)
  if i==0:
   fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8EEF2');props.append(fill)
  for p in cell.paragraphs:
   p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
   if j==1:p.alignment=1
   for run in p.runs:run.font.size=Pt(10);run.bold=i==0
doc.add_paragraph('작업용 장갑의1회는 점령 성공1회를 뜻한다. 발동한 효과의 지속 시간과 아이템 사용 횟수는 별개다.')
doc.add_page_break();doc.add_heading('사용과 시작 배치 규칙',1)
for title,text in c['rules']:
 p=doc.add_paragraph();p.add_run(title+'  ').bold=True;p.add_run(text)
doc.add_paragraph(c['basis'])
doc.add_paragraph(c['verification'])
doc.add_paragraph().add_run().add_picture(str(D/'images'/c['capture'][0]),width=Inches(5.8))
doc.add_paragraph(c['capture'][1],'Caption')
doc.save(D/'ItemUsageReport_20261018.docx');shutil.copy2(D/'ItemUsageReport_20261018.docx',O/'CurrentItemUsage.docx')
table=['| 아이템 | 사용량 | 효과 | 소모 시점 |','|---|---|---|---|']+['| '+' | '.join(v.replace('\n',' ') for v in row)+' |' for row in c['rows']]
md=['# '+c['title'],c['date'],c['intro'],'## 아이템별 획득당 사용량','\n'.join(table)]
md += ['## 소모 규칙']+[f'**{title}** {text}' for title,text in c['rules']]+['## 적용 근거',c['basis'],'## 검증과 적용 범위',c['verification'],c['source']]
markdown='\n\n'.join(md)+'\n';(D/'ItemUsageReport_20261018.md').write_text(markdown,encoding='utf-8');(O/'CurrentItemUsage.md').write_text(markdown,encoding='utf-8')
parts=['<h1>'+c['title']+'</h1><p>'+c['date']+'</p><p>'+c['intro']+'</p><table><thead><tr>'+''.join('<th>'+x+'</th>' for x in ['아이템','사용량','효과','소모 시점'])+'</tr></thead><tbody>']
for row in c['rows']:parts.append('<tr>'+''.join('<td>'+html.escape(v).replace('\n','<br>')+'</td>' for v in row)+'</tr>')
parts.append('</tbody></table><h2>소모 규칙</h2>')
for title,text in c['rules']:parts.append('<p><strong>'+title+'</strong> '+html.escape(text)+'</p>')
parts += ['<h2>적용 근거</h2><p>'+c['basis']+'</p>','<h2>검증과 적용 범위</h2><p>'+c['verification']+'</p>']
image=base64.b64encode((D/'images'/c['capture'][0]).read_bytes()).decode()
parts.append('<figure><img src="data:image/png;base64,'+image+'"><figcaption>'+c['capture'][1]+'</figcaption></figure>')
page='<!doctype html><html lang="ko"><meta charset="utf-8"><title>'+c['title']+'</title><style>body{max-width:1000px;margin:40px auto;padding:0 24px;font:16px/1.65 "Malgun Gothic",sans-serif;color:#172c3c}table{border-collapse:collapse;width:100%}td,th{border:1px solid #d9d9d9;padding:10px}th{background:#e8eef2}img{width:100%}figcaption{font-size:14px}</style><body>'+''.join(parts)+'</body></html>'
(D/'ItemUsageReport_20261018.html').write_text(page,encoding='utf-8');(O/'CurrentItemUsage.html').write_text(page,encoding='utf-8')
print('Current9 item report: DOCX/MD/HTML generated.')
