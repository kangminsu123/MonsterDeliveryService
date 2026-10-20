from docx.shared import Inches

def append_case06(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 for title,key in [('문제와 기대 동작','problem'),('원인','cause'),('수정 방법','method'),('해결 결과','result')]:
  p=doc.add_paragraph();p.add_run(title+'  ').bold=True;p.add_run(c[key])
 doc.add_heading('확정한 획득당 사용량',1)
 t=doc.add_table(rows=1,cols=2);t.style='Table Grid'
 for cell,text in zip(t.rows[0].cells,['현재 아이템','사용량']):cell.text=text
 for name,amount in [('복싱 글러브와 압축 뚫어뻥','각3회'),('뿅망치 슬라임 아이스크림 만년필','각2회'),('러닝화와 에어캡','각1회'),('작업용 장갑','점령 성공1회')]:
  for cell,text in zip(t.add_row().cells,[name,amount]):cell.text=text
 doc.add_paragraph('아이스크림2회는 교체 후 코드의 기본값이다. 옛 선풍기의4초 배터리는 적용하지 않는다. 현행 상세표는 별도 ItemUsageReport_20261018.docx를 참고한다.')
 doc.add_page_break();doc.add_heading('실제 남은 횟수 표시와 검증',1)
 name,caption=c['images'][1];doc.add_paragraph().add_run().add_picture(str(D/'images'/name),width=Inches(6.1));doc.add_paragraph(caption,'Caption')
 for title,key in [('검증 범위','verification'),('재시도','retry'),('적용 한계','limits'),('변경 파일','files')]:
  p=doc.add_paragraph();p.add_run(title+'  ').bold=True;p.add_run(c[key])
