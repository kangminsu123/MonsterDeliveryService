from docx.shared import Inches

def append_case07(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 for title,key in [('요청과 기대 동작','problem'),('원인','cause'),('수정 방법','method'),('해결 결과','result'),('문서 정리','cleanup')]:
  p=doc.add_paragraph();p.add_run(title+'  ').bold=True;p.add_run(c[key])
 doc.add_page_break();doc.add_heading('실제 상자 후보와 활성 위치',1)
 for name,caption in c['images'][:2]:
  doc.add_paragraph().add_run().add_picture(str(D/'images'/name),width=Inches(6.15));doc.add_paragraph(caption,'Caption')
 doc.add_page_break();doc.add_heading('맵 네 모서리에서의 실제 시작 화면',1)
 t=doc.add_table(rows=2,cols=2);t.autofit=False
 for cell,(name,caption) in zip([cell for row in t.rows for cell in row.cells],c['images'][2:]):
  cell.paragraphs[0].add_run().add_picture(str(D/'images'/name),width=Inches(3.03));cell.add_paragraph(caption,'Caption')
 for title,key in [('검증','verification'),('재시도와 원인 확인','retry'),('적용 범위와 한계','limits')]:
  p=doc.add_paragraph();p.add_run(title+'  ').bold=True;p.add_run(c[key])
 doc.add_paragraph(c['files'])
