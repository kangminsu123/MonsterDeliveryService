from docx.shared import Inches
def append_case09(doc,D,c):
 doc.add_page_break();doc.add_heading(c['title'],1);doc.add_paragraph(c['date'])
 for title,key in [('문제와 기대 동작','problem'),('원인','cause'),('수정 방법','method'),('해결 결과','result'),('검증','verification'),('재시도','retry'),('한계','limits'),('변경 파일','files')]:
  p=doc.add_paragraph();p.add_run(title+'  ').bold=True;p.add_run(c[key])
 for title,images in [('강화 망치의 실제 연결 화면',c['images'][:2]),('강화 아이스크림의 실제 연결 화면',c['images'][2:])]:
  doc.add_page_break();doc.add_heading(title,1);doc.add_paragraph(c['capture_note'])
  for name,caption in images:
   doc.add_paragraph().add_run().add_picture(str(D/'images'/name),width=Inches(6.0));doc.add_paragraph(caption,'Caption')
