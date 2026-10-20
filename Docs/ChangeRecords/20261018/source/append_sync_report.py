from docx.shared import Inches
import json

def append_case04(doc,D,c):
 def p(text):
  paragraph=doc.add_paragraph(text);paragraph.paragraph_format.keep_together=True
 def section(title,key):doc.add_heading(title,1);p(c[key])
 doc.add_page_break();doc.add_heading(c['title'],1);p(c['date'])
 for title,key in [('문제와 점검 범위','problem'),('원인','cause'),('수정한 방법','method'),('결과','result')]:section(title,key)
 doc.add_page_break();doc.add_heading('뿅망치 수정 전후와 실제 진단 화면',1);p(c['capture_note'])
 t=doc.add_table(rows=1,cols=2);t.autofit=False
 for cell,(name,caption) in zip(t.rows[0].cells,c['images'][:2]):
  cell.paragraphs[0].add_run().add_picture(str(D/'images'/name),width=Inches(3.05));cell.add_paragraph(caption,'Caption')
 doc.add_paragraph().add_run().add_picture(str(D/'images'/c['images'][2][0]),width=Inches(5.6));doc.add_paragraph(c['images'][2][1],'Caption')
 p('기본4상태는 본인과 상대에서 각각 IDLE WALK IDLE_ARMED WALK_ARMED 전환을 확인했다. 작업 장갑은 패시브 양손 장비이며 운동화·에어캡은 손 모델 없이 상태 이펙트로 표시한다. 별도 신규 점프·피격 클립은 추가하지 않았다.')
 doc.add_page_break();doc.add_heading('동일 지연 조건의 전후 실측',1);p(c['measurement'])
 values=[('수신 주기 평균','18.33Hz','23.97Hz'),('방장 동기화 처리 평균','0.1971ms','0.1207ms'),('참가자 동기화 처리 평균','0.2521ms','0.2353ms'),('입력 승인 평균','162.60ms','162.72ms'),('참가자 프레임 평균','16.77ms','16.75ms'),('7초 이동 프레임P95','17.90ms','17.55ms'),('방장 송신량 평균','10.63KB/s','12.42KB/s'),('참가자 수신량 평균','11.03KB/s','12.91KB/s'),('4명 샘플JSON 크기','2586B','2137B'),('4명 샘플 압축 크기','536B','491B')]
 t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
 for cell,text in zip(t.rows[0].cells,['측정 항목','최적화 전','최적화 후']):cell.text=text
 for row in values:
  for cell,text in zip(t.add_row().cells,row):cell.text=text
 p(c['interpretation'])
 doc.add_page_break()
 for title,key in [('검증과 재실행','verification'),('재시도와 진단','retry'),('남은 범위와 제한','limits'),('변경 파일','files')]:section(title,key)
 p('원본 evidence/sync-comparison.json · sync-verification.json · sync-audit-after.txt\n실제 양측 sync-connected-after-host/client.txt 및 sync-movement-host/client.txt\n이전 보고서와 이번 변경 전후 source/case04 · 이번 차이 evidence/sync-session.diff')
