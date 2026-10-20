import sys
sys.dont_write_bytecode = True
from pathlib import Path
import json, html, base64, csv
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

D = Path(__file__).resolve().parent.parent
data_path = D/'source/report_data.json'
data = json.loads(data_path.read_text(encoding='utf-8'))
doc = Document(D/'source/report_template.docx')
for child in list(doc._element.body):
    if child.tag != qn('w:sectPr'): doc._element.body.remove(child)
for section in doc.sections:
    section.top_margin=Inches(.65);section.bottom_margin=Inches(.65)
    section.left_margin=Inches(.7);section.right_margin=Inches(.7)
    for area in (section.header,section.footer):
        for p in area.paragraphs:p.text=''
for style_name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    s=doc.styles[style_name];s.font.name='맑은 고딕';s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Caption'].font.size=Pt(9)
def para(t):doc.add_paragraph(t)
def section(title,text):doc.add_heading(title,1);para(text)
def pic(name,caption,w=6.5):
    doc.add_paragraph().add_run().add_picture(str(D/'images'/name),width=Inches(w))
    doc.add_paragraph(caption,'Caption')
doc.add_paragraph(data['title'],'Title');para(data['date'])
para(data['summary'])
for title,key in [('문제와 기대 동작','problem'),('확인한 원인','cause'),('수정한 방법','method'),('수정 후 해결 결과','result')]:section(title,data[key])
doc.add_page_break();section('실행 검증과 적용 범위',data['verification'])
para('게임 코드: Assets/Gameplay/DeliveryHeldItem.cs · FirstPerson()\n검사 코드: Assets/Editor/GloveRotationCheck.cs\n이전 관련 기록: Docs/ChangeRecords/20261016/README.md 사례02 (온라인 아이템 표시 경로)')
section('재실행 메뉴와 근거',data['evidence'])
section('남은 확인과 적용 범위',data['limits'])
doc.add_page_break();doc.add_heading('같은 조건의 수정 전후',1)
para(data['capture_note'])
for kind,label in [('BoxingGloves','복싱 글러브'),('WorkGloves','작업용 장갑')]:
    doc.add_heading(label,2)
    t=doc.add_table(rows=1,cols=2);t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for cell,(name,caption) in zip(t.rows[0].cells,[(f'20261017_{kind}_before_call12.png','수정 전 · 반복 호출 12회 후'),(f'20261017_{kind}_after_final.png','수정 후 · 반복 및 복귀 검사 후')]):
        cell.width=Inches(3.25)
        cell.paragraphs[0].add_run().add_picture(str(D/'images'/name),width=Inches(3.1))
        cell.add_paragraph(caption,'Caption')
    para('처음 보이는 자세는 유지하면서, 이후 갱신으로 장갑이 뒤집히는 현상을 제거했다. 각도 측정은 600회 전체의 최댓값이며 정지 화면만으로 시간 경과를 판정하지 않았다.')
doc.add_page_break();doc.add_heading('실제 게임 화면과 확인 범위',1)
para('2026-10-18 SampleScene Play Mode. 로컬 연결 상태를 설정해 실제 온라인 표시 함수를 호출한 검수용 재현이다. 화면은 Unity Game View 원본이며 다른 PC와의 실제 접속 화면은 아니다.')
pic('20261018_BoxingGloves_play_after.png','복싱 글러브 · 실제 표시 함수 180프레임 후 회전 변화 0.00000도',6.1)
pic('20261018_WorkGloves_play_after.png','작업용 장갑 · 실제 표시 함수 180프레임 후 회전 변화 0.00000도',6.1)
for p in doc.paragraphs:
    p.paragraph_format.keep_together=True
doc.save(D/'ChangeReport_20261017.docx')

parts=[f'<h1>{html.escape(data["title"])}</h1><p>{html.escape(data["date"])}</p>']
md=[f'# {data["title"]}',data['date'],'관련 요청: R01, R02, R03 · 사례01 · 버그 수정 · 상태: Editor/Play 재현 해결 완료']
for title,key in [('요약','summary'),('문제와 기대 동작','problem'),('원인','cause'),('사용한 방법','method'),('해결 결과','result'),('확인 근거','verification'),('캡처 설명','capture_note'),('재실행과 로그','evidence'),('남은 제한','limits')]:
    text=data[key];parts.append(f'<h2>{title}</h2><p>{html.escape(text).replace(chr(10),"<br>")}</p>');md.extend([f'## {title}',text])
for name,caption in data['images']:
    parts.append(f'<figure><img src="images/{name}" alt="{html.escape(caption)}"><figcaption>{html.escape(caption)}</figcaption></figure>')
    md.append(f'![{caption}](images/{name})')
style='<meta charset="utf-8"><style>body{max-width:1040px;margin:45px auto;padding:0 24px;font:16px/1.7 "Malgun Gothic",sans-serif;color:#172c3c}h1{font-size:30px}h2{margin-top:30px}img{width:100%;height:auto}figure{margin:26px 0}figcaption{color:#526574;font-size:14px}</style>'
content='<!doctype html><html lang="ko"><head>'+style+'<title>'+data['title']+'</title></head><body>'+''.join(parts)+'</body></html>'
(D/'index.html').write_text(content,encoding='utf-8')
(D/'README.md').write_text('\n\n'.join(md)+'\n',encoding='utf-8')
with (D/'requests.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f);w.writerow(['요청 ID','대상','요청 요약','관련 사례','처리 상태']);w.writerows(data['requests'])
out=Path(sys.argv[1]) if len(sys.argv)>1 else None
if out:
    out.mkdir(parents=True,exist_ok=True)
    import shutil
    shutil.copy2(D/'ChangeReport_20261017.docx',out/'GloveRotationFix.docx')
    for name,_ in data['images']:
        encoded=base64.b64encode((D/'images'/name).read_bytes()).decode()
        content=content.replace('images/'+name,'data:image/png;base64,'+encoded)
    (out/'GloveRotationFix.html').write_text(content,encoding='utf-8')
