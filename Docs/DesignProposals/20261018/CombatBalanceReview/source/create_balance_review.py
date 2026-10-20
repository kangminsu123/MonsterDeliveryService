from pathlib import Path
import json, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

W=Path(r'C:/Users/xlfhf/Documents/Codex/2026-10-17/dbslx')
R=Path(r'C:/Users/xlfhf/Documents/Codex/2026-09-18/https-github-com-kangminsu123-monsterdeliveryservice-git/MonsterDeliveryService')
D=R/'Docs/DesignProposals/20261018/CombatBalanceReview';D.mkdir(parents=True,exist_ok=True)
O=W/'outputs';O.mkdir(exist_ok=True)
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
sec.header_distance=sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='맑은 고딕';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
 st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=Pt(15)
 if name=='Normal':st.font.size=Pt(11)
 if name=='Title':st.font.size=Pt(25);st.font.bold=True;st.paragraph_format.line_spacing=Pt(32)
 if name=='Heading 1':st.font.size=Pt(18);st.font.bold=True;st.paragraph_format.space_before=Pt(5);st.paragraph_format.line_spacing=Pt(24)
 if name=='Heading 2':st.font.size=Pt(12);st.font.bold=True;st.paragraph_format.space_before=Pt(9)
 if name=='Subtitle':st.font.size=Pt(10.5)
 # Strip theme title decorations and any paragraph border.
 ppr=st.element.find(qn('w:pPr'))
 if ppr is not None:
  for el in list(ppr):
   if el.tag==qn('w:pBdr'):ppr.remove(el)
 st.font.underline=False

md=[]
def p(text,boldlead=None):
 q=doc.add_paragraph()
 if boldlead and text.startswith(boldlead):q.add_run(boldlead).bold=True;q.add_run(text[len(boldlead):])
 else:q.add_run(text)
 md.append(text+'\n')
 return q
def h(text,level=2):doc.add_heading(text,level);md.append('#'*(level+1)+' '+text+'\n')
def page(title):
 h(title,1);doc.paragraphs[-1].paragraph_format.page_break_before=True
def table(headers,rows,widths,size=10.5):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=v
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
   pr.append(b);m=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');m.append(el)
   pr.append(m)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2' if i==0 else 'FFFFFF');pr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=Pt(13.5)
    if j==0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.size=Pt(size);run.bold=(i==0)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(row)+' |' for row in rows)+'\n')


doc.add_paragraph('아이템 전투 강화 수치 검토 보고서','Title')
doc.add_paragraph('Monster Delivery Service  2026년 10월 18일  적용 전 검토안','Subtitle')
md.append('# 아이템 전투 강화 수치 검토 보고서\n\n2026년 10월 18일 · 적용 전 검토안\n')
p('기존 비살상 아이템 전투의 영향력을 높이는 상승 높이 보완을 포함한 일곱 가지 항목을 현재 코드와 비교했다. 상대를 더 멀리 밀고 오래 묶으며, 더 먼 거리에서 당기거나 빠르게 빙결시키는 방향이다. 이 보고서는 요청 수치와 예상 영향을 검토하기 위한 자료다.')
p('현재 게임에는 적용하지 않았다. 변경안의 이동 거리·비행 시간은 코드에 근거한 단순 계산이며 이번 Unity 실측 결과가 아니다. 기절·빙결 면역을 5초로 늘리자는 앞선 제안도 채택하지 않은 상태로 구분한다.',boldlead='현재 게임에는 적용하지 않았다.')
h('요청 수치 비교')
table(['대상과 항목','현재 기준','요청 변경안','변화'],[
 ['복싱 글러브 수평 넉백','약 16.2m 계산 기준','약 32.4m 목표','거리 2배'],
 ['복싱 글러브 상승 높이','약 2.7m 계산 기준','약 5.4m 검토 목표','높이 2배 제안'],
 ['뿅망치 기절 지속','1.2초','5초','약 4.17배'],
 ['압축 뚫어뻥 사거리','15m','30m','2배'],
 ['아이스크림 발사속도','조준 방향 18m/s','조준 방향 36m/s','2배'],
 ['아이스크림 빙결 지속','1초','3초','3배'],
 ['슬라임 획득당 사용량','2회','3회','50퍼센트 증가']
],[2.0,1.7,1.95,1.45])
p('아이스크림의 “기절 3초” 요청은 기존 빙결이 이동·행동·점령을 막는 시간을 3초로 늘리는 뜻으로 정리한다. 별도 기절 3초를 추가해 총 6초로 만드는 안이 아니다.')
h('유지하는 현행 규칙')
p('복싱 글러브는 획득당 2회, 슬라임만 변경안에서 3회이며 나머지 7종은 각 1회를 유지한다. 작업용 장갑의 1회는 점령 성공 1회다. 기본 점령 3초와 장갑 점령 1.5초, 우체통 재점령 보호 20초 및 남은 시간 표시는 유지한다.')
p('공격 간격은 복싱 0.6초, 뿅망치 0.75초, 뚫어뻥 0.7초, 아이스크림 0.6초를 유지하는 비교안이다. 기절·빙결이 끝난 뒤 공통 면역 2초도 유지한다. 지속시간 증가와 쿨타임 증가는 서로 다른 변경이다.')
p('빈 보조 슬롯 3칸 시작, 기존 9종 균등 획득, 상자 100개 후보와 활성 50개, 네 모서리 랜덤 시작을 유지한다. 새 무기·체력·사망·리스폰은 이번 안에 포함하지 않는다.')

page('변경안에서 예상되는 전투 흐름')
h('복싱 글러브는 멀리 밀면서 위로도 크게 띄우기')
p('현재 수평 밀치기 속도는 약 31.18m/s이고 초당 30m/s씩 줄어든다. 장애물 없는 첫 타격의 단순 수평 거리 계산은 약 16.2m다. 목표 거리를 32.4m로 잡으면 초기 수평 속도는 약 44.09m/s로 조정하는 해석이 맞다. 속도를 그대로 2배로 하면 약 64.8m로 계산되어 요청한 거리 2배를 넘어선다.')
p('후속 요청에 따라 수평 거리뿐 아니라 위로 띄우는 높이도 크게 늘린다. 수직 초기 속도 약 10.39m/s와 중력 가속도 20m/s²의 단순 계산상 현재 상승량은 약 2.7m다. 상승량도 2배인 약 5.4m를 우선 검토 목표로 제안하며, 필요한 수직 초기 속도는 약 14.70m/s다. 높이 2배는 수치가 지정되지 않은 요청에 대한 제안으로 확정값은 아니다.')
p('상승량은 맞은 위치부터 최고점까지의 추가 높이이며 지면에서의 절대 높이가 아니다. 평지에서 같은 높이로 돌아오는 단순 체공 시간은 약 1.04초에서 1.47초로 늘어난다. 공격 사거리 5m는 유지한다. 벽·천장·지붕·경사·절벽 복귀와 반복 타격 저항 때문에 실제 궤적과 거리는 달라진다.')
h('기절과 빙결로 점령할 시간 확보')
table(['행동 불가 효과','효과 시간','기본 점령 3초 대비','장갑 점령 1.5초 대비'],[
 ['뿅망치 기절','5초','계산상 2초 여유','계산상 3.5초 여유'],
 ['아이스크림 빙결','3초','이동 전환 여유 없음','계산상 1.5초 여유']
],[1.65,1.0,2.2,2.25])
p('위 비교는 적중 직후 이미 우체통을 바라보며 즉시 점령할 수 있을 때의 계산이다. 접근·무기 동작·전환·네트워크 지연은 제외했다. 보호 중인 우체통은 상태 이상과 관계없이 남은 20초 제한이 끝나야 점령할 수 있다.')
p('뿅망치 5초는 이번 안에서 행동 제한 증가가 가장 크다. 공격 성공이 점령으로 이어지기 쉬워지는 대신 맞은 사람의 대기와 집중 공격 부담이 커질 수 있다. 종료 후 면역 2초를 유지하면 같은 대상에게 다음 기절·빙결이 적용 가능한 가장 이른 시점은 뿅망치 적중 약 7초 뒤, 아이스크림 적중 약 5초 뒤다. 넉백·당기기·감속까지 모두 막는 면역은 아니다.')
h('뚫어뻥은 먼 상대를 전장 안으로 끌어오기')
p('사거리는 30m로 늘리고 출발 속도 36m/s와 복귀 속도 44m/s는 유지하는 안이다. 최대 거리의 단순 왕복은 약 0.76초에서 1.52초로 늘어난다. 준비 시간·손 위치·충돌·상대 이동을 제외한 계산이다. 더 먼 상대를 끌어올 수 있지만 날아갔다 돌아오는 동안 다음 사용을 못 하는 시간도 길어진다.')
h('아이스크림과 슬라임의 압박 증가')
p('아이스크림은 조준 방향 속도를 18에서 36m/s로 올리고 추가 상승 속도 3m/s와 중력을 유지하는 해석이다. 속도 벡터 전체를 두 배로 해 상승도 6m/s로 만드는 것은 별도 변경이다. 거리 10m의 단순 비행 시간은 약 0.56초에서 0.28초로 줄어든다. 투척 준비 0.22초와 낙하·충돌은 이 계산에서 제외한다.')
p('슬라임은 한 번 획득해 세 번 던질 수 있어 진입로·우체통 주변을 더 오래 압박할 수 있다. 직접 감속 3초와 40퍼센트 감소, 바닥 반경 3m와 유지 5초는 유지한다. 세 번 던져도 감속 비율을 세 배로 중첩하는 안은 아니다.')

page('적용 검토와 후속 확인 기준')
p('요청한 수치 그대로를 강화안으로 비교하되, 공격이 유리해진 정도와 맞는 사람이 대응할 수 있는지를 함께 본다. 지금 단계에서는 수치를 낮추거나 다른 규칙을 추가하지 않는다.')
h('구현할 경우 함께 맞출 곳')
p('현재 로컬과 멀티 서버에 복싱 충격량, 뿅망치 기절 시간, 뚫어뻥 거리, 아이스크림 발사속도가 나뉘어 기록되어 있다. 한쪽 상수만 바꾸면 서로 다른 효과가 생길 수 있다. 두 경로의 값을 함께 맞추고 빙결 화면·피드백 문구와 사용량 HUD까지 같은 기준을 사용해야 한다.')
p('슬라임 사용량은 공통 Ammo 기준으로 3회가 되고, 획득·소모·슬롯 소진·서버 동기화도 이를 따라야 한다. 뚫어뻥 장거리 당기기는 기존 당기기 종료 시간 안에 실제로 복귀하는지, 아이스크림은 빠른 이동에서도 작은 상대나 벽을 건너뛰지 않는지 확인해야 한다.')
h('후속 플레이에서 비교할 항목')
table(['대상','확인할 동작'],[
 ['복싱 글러브','수평 거리 약 2배와 상승량 증가. 최고점·체공·지붕/천장·절벽 복귀·반복 저항과 양측 표시.'],
 ['뿅망치','기절 5초와 종료 후 면역 2초. 효과 종료 전 조작 복귀나 조기 점령 재개 없음.'],
 ['뚫어뻥','15m 이후부터 30m까지 적중·당기기, 벽 차단, 끝점 복귀, 비행 중 손 모델 숨김.'],
 ['아이스크림','발사속도와 3초 빙결 일치. 벽 앞 발사·이동 표적·빙결 종료·에어캡·공통 면역.'],
 ['슬라임','획득 3회 표시, 세 번만 소모, 마지막 투척 유지, 로컬과 참가자 횟수 일치.'],
 ['전체 경기','공격 후 탈환 성공률, 행동 불가 체감, 장비별 사용 가치, 지연 상황에서 대응 가능성.']
],[1.25,5.85])
h('보고서의 범위와 남은 결정')
p('이 문서는 요청 수치 비교와 코드 검토만 완료한 보고서다. 새 수치로 Unity를 실행하거나 별도 멀티 플레이를 검증하지 않았다. 예상 거리는 수평 감속 모형, 비행 시간은 거리와 속도에 따른 근사값이다. 강화 후 실제 체감과 밸런스는 아직 미확인이다.')
p('구현 전에 확인할 부분은 복싱 수평 거리 2배와 상승량 2배 제안, 아이스크림 조준 방향 36m/s와 상승 보정 3m/s 유지다. 복싱은 기존 높이 유지 해석을 폐기하고 더 높이 띄우는 요청을 반영했다. 요청한 뿅망치 5초·빙결 3초·슬라임 3회는 임의로 낮추지 않고 그대로 정리했다.')
h('현재 값의 확인 근거')
p('GloveCombat.cs와 DeliveryItemState.cs에서 수평·수직 충격량과 감속을, MonsterDeliveryPrototype.cs와 DeliveryNetworkPrediction.cs에서 중력 20m/s²를 확인했다. ToyHammerCombat.cs는 기절 1.2초, PlungerCombat.cs는 사거리 15m, IceCreamCombat.cs는 속도 18m/s와 빙결 1초다. MonsterDeliveryPrototype.cs의 Ammo는 슬라임 2회이며, DeliveryNetworkGameplay.cs에서 별도 서버 처리 값을 대조했다.')
p('현행 기준은 사례08의 재점령 보호와 최신 아이템 사용 문서다. 이 보고서가 현행 사용 기준을 대체하지 않으며, 폐기한 기본 무기·체력·사망·리스폰 기획안도 되살리지 않는다.')

header=sec.header.paragraphs[0];header.text='MONSTER DELIVERY SERVICE';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('아이템 전투 강화 검토안  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for run in footer.runs:run.font.size=Pt(9)
for para in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
 grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');para._p.get_or_add_pPr().append(grid)
doc.core_properties.title='아이템 전투 강화 수치 검토 보고서';doc.core_properties.subject='미적용 요청 수치 비교';doc.core_properties.language='ko-KR'
target=D/'CombatBalanceReview_20261018.docx';doc.save(target)
markdown='\n'.join(md).replace('\\n','\n');(D/'CombatBalanceReview_20261018.md').write_text(markdown,encoding='utf-8')
shutil.copy2(target,O/'CombatBalanceReview.docx');(O/'CombatBalanceReview.md').write_text(markdown,encoding='utf-8')
print(target)
