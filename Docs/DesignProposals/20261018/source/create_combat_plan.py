from pathlib import Path
import json, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

W=Path(r'C:/Users/xlfhf/Documents/Codex/2026-10-17/dbslx')
R=Path(r'C:/Users/xlfhf/Documents/Codex/2026-09-18/https-github-com-kangminsu123-monsterdeliveryservice-git/MonsterDeliveryService')
D=R/'Docs/DesignProposals/20261018';D.mkdir(parents=True,exist_ok=True)
O=W/'outputs';O.mkdir(exist_ok=True)
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
sec.header_distance=sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='맑은 고딕';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
 st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=Pt(15)
 if name=='Normal':st.font.size=Pt(11)
 if name=='Title':st.font.size=Pt(25);st.font.bold=True;st.paragraph_format.line_spacing=Pt(32)
 if name=='Heading 1':st.font.size=Pt(18);st.font.bold=True;st.paragraph_format.space_before=Pt(5);st.paragraph_format.line_spacing=Pt(24)
 if name=='Heading 2':st.font.size=Pt(12);st.font.bold=True;st.paragraph_format.space_before=Pt(9)
 if name=='Subtitle':st.font.size=Pt(10.5)
 # Strip theme title decorations and any paragraph border.
 ppr=st.element.find(qn('w:pPr'))
 if ppr is not None:
  for el in list(ppr):
   if el.tag==qn('w:pBdr'):ppr.remove(el)
 st.font.underline=False

md=[]
def p(text,boldlead=None):
 q=doc.add_paragraph()
 if boldlead and text.startswith(boldlead):q.add_run(boldlead).bold=True;q.add_run(text[len(boldlead):])
 else:q.add_run(text)
 md.append(text+'\n')
 return q
def h(text,level=2):doc.add_heading(text,level);md.append('#'*(level+1)+' '+text+'\n')
def page(title):
 h(title,1);doc.paragraphs[-1].paragraph_format.page_break_before=True
def table(headers,rows,widths,size=10.5):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=v
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
   pr.append(b);m=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');m.append(el)
   pr.append(m)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2' if i==0 else 'FFFFFF');pr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=Pt(13.5)
    if j==0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.size=Pt(size);run.bold=(i==0)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(row)+' |' for row in rows)+'\n')

doc.add_paragraph('기본 무기와 체력 전투 확장 기획안','Title')
doc.add_paragraph('Monster Delivery Service  2026년 10월 18일  구현 전 검토 초안','Subtitle')
md.append('# 기본 무기와 체력 전투 확장 기획안\n\n2026년 10월 18일 · 구현 전 검토 초안\n')
p('플레이어가 항상 기본 무기로 싸울 수 있고, 파밍한 보조 아이템으로 전투와 점령을 유리하게 만드는 방향을 제안한다. 체력이 0이 되면 사망한 뒤 리스폰하며, 상대를 쓰러뜨려 확보한 시간을 우체통 점령에 활용한다.')
p('이번 요청은 문서 작성까지다. 아래 전투 규칙은 아직 게임에 적용되지 않았다. 기본 무기·보조 아이템 구조, 체력·사망·리스폰, 원거리·고위력 근접·방어막 추가는 요청 방향이다. 이름·피해량·시간·키 배치는 모두 검토용 제안이다.',boldlead='이번 요청은 문서 작성까지다.')
h('전투를 추가하는 이유')
p('현재는 아이템이 없으면 공격할 수 없고, 아이템을 소진하면 상대를 피하는 편이 유리할 수 있다. 기본 무기로 전투를 지속하고 사망으로 짧은 이탈 시간을 만들어, 공격 성공이 점령 기회로 연결되도록 한다. 실제로 전투가 유리해지는지는 다인 플레이에서 확인한다.')
h('첫 테스트의 기본 규칙')
table(['항목','제안 기준'],[
 ['기본 무기','맨손 펀치 무제한 보유. 보조 슬롯을 차지하지 않는다.'],
 ['체력과 기본 피해','최대 체력 100. 펀치 피해 25, 사거리 2.2m, 공격 간격 0.6초.'],
 ['기본 공격 제약','벽을 관통하지 않고 가장 가까운 유효 상대 1명만 타격. 기절·강한 넉백 없음.'],
 ['보조 장비','기존 3칸 유지. 기본 무기와 보조 장비 중 하나만 손에 선택한다.'],
 ['승패','우체통 10개와 종료 시 점령 수 비교를 우선 유지. 처치 점수는 승패에 가산하지 않는다.'],
 ['시작 배치','기존 네 모서리 랜덤 시작, 상자 100개 후보와 활성 50개 유지.'],
 ['사망과 복귀','체력 0이면 행동 중단. 약 5초 뒤 체력 100과 기본 무기로 복귀.']
],[1.35,5.75])
p('흐름은 상자 파밍 → 보조 아이템으로 유리한 상황 만들기 → 기본 무기로 전투 → 상대 이탈 중 점령 → 재교전이다. 활성 우체통 3개나 보유 시간 점수 방식은 이번 초안에 포함하지 않는다.')

page('추가 아이템과 기존 장비의 역할')
p('원거리·고위력 근접·방어막 3종을 새 보조 아이템으로 제안한다. 회복 아이템은 “체력 추가”에 연결한 선택 제안이며 추가 여부를 따로 결정한다. 아래 이름과 수치는 확정값이 아니다.')
table(['추가 장비 가칭','효과와 초기 수치','횟수'],[
 ['고무줄 발사기','직선 투사체 1발. 피해 35, 속도 30m/s, 최대 거리 25m. 단일 적중, 관통·기절 없음.','1회'],
 ['대형 택배 망치','근접 2.5m 안 상대 1명에게 피해 60. 준비 0.45초, 재사용 1.2초. 빗나가도 소모, 기절 없음.','1회'],
 ['완충 방어막','5초 동안 체력 피해를 합계 40까지 흡수. 남은 피해는 체력에 적용. 행동 방해는 차단하지 않는다.','1회'],
 ['응급 구급팩 선택안','1초 사용 후 체력 40 회복. 최대 체력 초과 불가. 이동·공격·피격 시 취소하고 소모하지 않는다.','1회']
],[1.25,5.15,.7])
h('기존 9종 유지 기준')
table(['기존 아이템','전투에서의 역할','횟수'],[
 ['복싱 글러브','기본 펀치의 강화 보조 장비. 타격 피해 25 제안, 기존 강한 넉백 유지.','2회'],
 ['압축 뚫어뻥','상대를 끌어내고 위치를 무너뜨림. 체력 피해 0.','1회'],
 ['뿅망치','1.2초 기절로 공격 기회 확보. 체력 피해 0.','1회'],
 ['슬라임','직접 감속과 바닥 영역으로 진입·도주 방해. 체력 피해 0.','2회'],
 ['아이스크림','1초 빙결로 행동 제한. 체력 피해 0.','1회'],
 ['러닝화','5초 이동 강화로 추격·도주·점령 이동 지원.','1회'],
 ['에어캡','기존 방어 대상인 행동 방해 공격 1회 차단. 순수 체력 피해는 막지 않는 안.','1회'],
 ['작업용 장갑','선택 착용 중 점령 1.5초. 점령 성공 시 소모.','성공 1회'],
 ['만년필','4초 시야 방해로 공격·탈환 기회 확보. 체력 피해 0.','1회']
],[1.25,5.15,.7],10)
p('기존 9종에 필수 추가 3종을 합치면 보조 장비는 12종이며, 구급팩까지 채택하면 13종이다. 기본 무기는 이 수에 포함하지 않는다. 획득 확률은 기존 각 1/9에서 재설계해야 하며 아직 결정하지 않는다.')

page('피해 사망과 리스폰 규칙')
h('체력 피해와 두 방어 아이템')
p('피해는 방어막 흡수량부터 차감한 뒤 남은 양만 체력에서 뺀다. 예를 들어 방어막 40이 남은 상태에서 피해 60을 받으면 방어막은 사라지고 체력은 20 줄어든다. 방어막은 회복되지 않으며 재사용 시 수치를 더하지 않고 40과 5초로 갱신하는 안이다.')
p('에어캡은 기절·빙결·넉백·당기기 등 기존 방어 대상의 행동 방해를 1회 막는다. 완충 방어막은 체력 피해를 막는다. 둘을 함께 켰을 때 복싱 타격을 받으면 에어캡이 넉백을 막고 방어막이 피해를 흡수하는 안이다. 슬라임 바닥 감속은 기존처럼 에어캡으로 막지 않는다. 두 아이템의 효과를 HUD와 설명에서 구분한다.')
h('체력이 0이 되는 순간')
p('서버가 사망을 한 번 확정한다. 이동·공격·상자 획득·점령을 중단하고 진행 중인 준비 동작과 점령 진행을 취소한다. 사망자는 충돌과 목표 판정에서 제외한다. 보조 슬롯과 강화·방어·상태 이상은 비우며, 아이템을 바닥에 떨어뜨리는 규칙은 넣지 않는 안이다.')
p('소유한 우체통은 사망만으로 중립화하지 않는다. 상대가 직접 탈환해야 한다. 이미 발사된 투사체와 생성된 슬라임 영역은 정상 수명이 끝날 때까지 유지해 동시 처치도 허용하는 안이다. 진행 중 뚫어뻥 왕복과 붙잡기는 종료해 죽은 사용자에게 끌려가지 않게 한다.')
h('다시 전장에 복귀')
p('리스폰 대기 5초 후 체력 100, 기본 무기 선택, 빈 보조 슬롯 3개, 스태미나 100으로 복귀한다. 자동 체력 회복은 우선 두지 않는다. 구급팩을 채택하지 않으면 체력은 리스폰 때만 회복되는 초기안이다.')
p('경기 시작의 네 모서리 배정은 유지한다. 사망 후에는 적과 떨어지고 몸이 끼이지 않는 별도 안전 복귀 지점을 사용하도록 제안한다. 지점 수와 적과의 최소 거리는 맵에서 검토해야 한다. 복귀 무적은 2초를 초기안으로 두고 공격을 시작하면 즉시 해제한다. 무적 중 점령과 아이템 수집은 막아 일방적인 목표 획득을 방지한다.')
h('조작과 화면')
p('좌클릭은 현재 선택한 무기를 사용한다. 숫자 1~3은 보조 슬롯 선택, Q는 기본 무기로 돌아가기, G는 선택한 보조 아이템 폐기로 제안한다. 장비 소진 후에는 기본 무기로 자동 복귀한다. 펀치와 보조 공격을 동시에 발동할 수 없으며 전환으로 공격 대기 시간을 없앨 수 없다.')
p('HUD에 현재 체력과 방어막을 별도로 표시한다. 피격 방향과 체력 감소, 사망 원인, 리스폰 남은 시간 및 복귀 무적을 안내한다. 회복 아이템을 채택하면 사용 진행과 취소도 표시한다. 이번 문서는 화면 시안이나 적용 완료 화면을 포함하지 않는다.')
h('전투가 점령으로 이어지는 예')
p('체력 100의 상대를 기본 펀치로 쓰러뜨리려면 4회 적중해야 한다. 고무줄 35와 기본 펀치 3회는 합계 110, 대형 망치 60과 기본 펀치 2회도 110이다. 방어막과 빗나감이 없을 때의 계산이다. 망치의 긴 준비와 소모 위험, 원거리의 명중 난도가 실제 교전에서 맞는지 비교한다.')

page('구현 순서와 검증할 항목')
p('현재 게임은 비살상 상태이며 체력·사망·리스폰이 구현되어 있지 않다. 이 초안은 기존 “체력·사망 없음” 기획을 대체할 다음 방향을 제안한다. 현재 아이템 사용 기준 문서는 현행 구현 자료로 유지하고 이 문서와 구분한다.')
h('단계별 적용')
p('1단계는 체력, 무제한 기본 무기, 사망·리스폰, 체력 HUD다. 기본 전투만으로 승패가 갈리는 시간과 복귀 시간을 확인한다. 기존 행동 방해 아이템은 공통 피격 처리와 연결하고, 사망자에게 추가 공격이 적용되지 않게 한다.')
p('2단계는 원거리 무기와 고위력 근접 무기, 피해 방어막이다. 같은 조준·벽 가림·소모 경로를 사용한다. 회복 아이템 채택 시 함께 넣고, 상자 획득 목록과 확률·설명·아이콘·슬롯 UI를 갱신한다.')
p('3단계는 멀티와 실제 경기 검증이다. 서버가 공격 승인·적중·방어막·체력·사망·리스폰·아이템 소모를 결정한다. 참가자에게 생존 상태, 체력, 방어막 잔량, 복귀 위치·방향·남은 시간과 무적 상태를 동기화한다. 중복 처치·중복 리스폰과 늦게 도착한 공격 입력을 방지하고 이전 프로토콜과의 호환 여부를 명시한다.')
h('직접 확인할 플레이 상황')
table(['상황','확인 기준'],[
 ['기본 전투','빈 보조 슬롯으로도 공격 가능. 한 공격의 중복 피해와 벽 뒤 적중 없음.'],
 ['장비와 방어','기존 9종 횟수 유지. 방어막 초과 피해와 에어캡 행동 방해 차단을 각각 확인.'],
 ['사망 순간','공격 준비·점령 취소, 기존 투사체 수명 유지, 사망자 재타격·움직임 차단.'],
 ['복귀와 회복','체력·슬롯·강화 초기화, 안전 지점과 무적 해제, 최대 체력과 회복 취소 검사.'],
 ['실제 멀티','별도 호스트·참가자에서 체력·죽음·복귀 일치. 지연·유실 환경의 적중과 중복 처리 확인.'],
 ['게임 흐름','교전 후 점령 성공률, 첫 교전 시간, 사망 후 복귀 시간, 코너별 승률과 장비 활용을 기록.']
],[1.2,5.9])
h('구현 전에 정할 사항')
p('기본 무기의 종류와 수치, 새 3종의 이름·수치, 구급팩 추가 여부, 에어캡과 피해 방어막의 역할 구분, 사망 시 장비 초기화, 리스폰 위치·무적, 새 획득 확률을 결정한다. 모든 제안 수치는 한 차례 플레이한 뒤 조정할 시작값이다.')
p('처치가 가능해져도 흩어진 우체통만 먹는 편이 계속 유리할 수 있다. 전투에서 이긴 뒤 실제 점령으로 이어지는지 먼저 확인하고, 이득이 약할 때 목표 배치나 점령 규칙을 추가 검토한다. 이번 초안만으로 전투 동기가 해결됐다고 판단하지 않는다.')
p('검토 기준은 2026년 10월 18일 사용자 대화와 현행 GAME_DESIGN, 아이템 사용 기준 및 사례07이다. 게임 코드·장면·수치·기존 사용 문서는 이번에 변경하지 않았다. 이 문서에는 Unity 구현 검사 결과를 기재하지 않는다.')

header=sec.header.paragraphs[0];header.text='MONSTER DELIVERY SERVICE';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('전투 확장 기획 초안  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for run in footer.runs:run.font.size=Pt(9)
doc.core_properties.title='기본 무기와 체력 전투 확장 기획안';doc.core_properties.subject='구현 전 검토 초안';doc.core_properties.language='ko-KR'
for para in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs]:
 grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');para._p.get_or_add_pPr().append(grid)
target=D/'CombatExpansionPlan_20261018.docx';doc.save(target)
markdown='\n'.join(md);(D/'CombatExpansionPlan_20261018.md').write_text(markdown,encoding='utf-8')
shutil.copy2(target,O/'CombatExpansionPlan.docx');(O/'CombatExpansionPlan.md').write_text(markdown,encoding='utf-8')
(D/'README.md').write_text('# 기본 무기와 체력 전투 확장 기획안\n\n2026-10-18 문서만 작성한 검토 초안이다. 게임 구현과 완료 기록이 아니다.\n\n[Word](CombatExpansionPlan_20261018.docx) · [편집 원문](CombatExpansionPlan_20261018.md)\n\n요청 범위: 기본 무기와 보조 슬롯, 체력·사망·리스폰, 원거리·고위력 근접·방어막 추가. 회복 아이템은 별도 선택 제안이다. 신규 이름과 모든 전투 수치는 미확정이다.\n',encoding='utf-8')
source=D/'source';source.mkdir(exist_ok=True);shutil.copy2(Path(__file__),source/'create_combat_plan.py')
print(target)
